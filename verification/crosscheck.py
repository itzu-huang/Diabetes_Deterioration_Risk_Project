# -*- coding: utf-8 -*-
"""
crosscheck.py —— 第 2 層:交付物數字「逐處」比對
================================================================
從報告 docx、簡報 pptx、海報 pptx、決策 App html 抽出全部文字單位(段落、表格格位、
文字框、講者備註、App 原始碼行),以「關鍵詞前後文＋數字」之樣式辨識每一處關鍵數字,
逐處與本次執行輸出(CSV)推導之新值比對。

〔稽核修正 B-03〕原版之判定為「六個文字來源中任一處出現新值即 ✓」,並以 reference_teacher
推導之「舊值」攔截漏改;34 項中 14 項新值＝舊值,這些項目任何一處改錯都攔不到(稽核實測:
報告 6 處「2.98」改為「2.89」、段落內 AUC 改值,crosscheck 仍 exit 0)。本版改為:
  (1) 每一處出現位置都必須等於本次輸出之新值(不論新舊值是否相同);
  (2) 以樣式辨識語境,例如成本倍數只認「長期風險低者之 x.xx 倍」「(A、B)…(C、D)之 x.xx 倍」
      「approximately x.xx times」與旁框註明「長期風險」之「x.xx 倍」等;
  (3) 納入 (A+B)/(C+D) 成本倍數(＝A、B 象限之每人平均成本 ÷ C、D 象限之每人平均成本,
      僅計首筆之 100 位病患;與報告 ¶699「長期風險高之兩象限（A、B）之餘生成本約為長期
      風險低者（C、D）之 2.98 倍」同義),以及交付物中全部 LSTM 數字(15/30/60 分鐘時界、
      逐人/合併、平均/中位數、彙總 RMSE 與 R²、族群落差、長度分組、偏相關…);
  (4) 每個不符印出:交付物、位置(報告 ¶N 或 表 x-y-z(T#)rRcC、簡報第 N 張本文/表/備註、
      海報文字框、App 行號)、前後文約 20 字、現值、應為值、來源欄位;並寫成
      crosscheck_mismatch.csv —— 學生重跑後,這份 CSV 就是「哪些數字要更新」的工作清單;
  (5) 某項在所有交付物中都找不到任何出現位置時亦列為失敗(樣式失效或數字已被刪除,須確認);
  (6) 期望值來自已捨入之 CSV 且恰在進位邊界(例:0.8875 → 0.887 或 0.888)時,列為 ▲ 邊界,
      兩個候選值皆暫不判錯,但提示須以全精度值確認(A-19 即此類);
  (7) 〔稽核修正 B-10〕共線性數字(報告 ¶521、¶574、¶585、¶802、¶803):相關係數取自
      output/eda_figures/correlation_matrix.csv;VIF、判定係數、拔靴少數方向次數、剔除特徵之
      AUC 取自 collinearity_check.py 之三個輸出(run_all.py --with-sensitivity 產生;原封包無此
      程式,故以原封包輸出執行時這些項目列為「無法檢查」)。¶521「前述各組之最大值(總膽固醇 9.3)」
      以「該五組變數之最大 VIF」比對(A-02:實為肌酸酐 15.8)。
  另保留兩項輔助檢查:App 常數區 const P 逐鍵比對;reference_teacher(另一次執行)之值若以
  「未被任何樣式涵蓋」之形式出現,列為疑似殘留(allowlist.csv 可豁免,但須寫明前後文)。

〔稽核修正 DOCS-17〕交付物與輸出目錄改以標記檔自動搜尋(不寫死資料夾名稱),或以
  --base/--out 指定。

[修正 26]〔數值由輸出推導、App 常數納入、負號正規化〕、[修正 31]〔檔名版本號〕之設計保留。

用法:
    python crosscheck.py                          # 自動尋找交付物與輸出目錄
    python crosscheck.py --base <交付物目錄> --out <程式與輸出目錄>
    python crosscheck.py --show-ok                # 另列出每一處相符之位置
    python crosscheck.py --show-uncovered         # 另列出同值但未被任何樣式涵蓋之出現處
    python crosscheck.py --mismatch-csv <路徑>    # 工作清單輸出位置(預設同目錄 crosscheck_mismatch.csv)
結束碼:0＝全部相符(邊界警告不計);1＝有不符、找不到出現處或無法檢查之項目。
"""

# ---------------------------------------------------------------------------
# [修正 32] Windows 主控台編碼防護。
#   本專案之輸出含 ✓ ≤ − ä ² ≈ 等字元,不在繁體中文 Windows 之預設編碼 cp950 內。
#   當 stdout 是「主控台」時 Python 走 WriteConsoleW,不受影響;但當 stdout 被
#   導向「管線」(例如 verify_all.py 以 capture_output=True 抓取子程序輸出,或
#   使用者自行 `python x.py > log.txt`)時,Python 改用地區編碼 cp950 編碼,
#   即拋出 UnicodeEncodeError 並中止 —— 程式本身沒錯,卻因為印不出一個勾勾而失敗。
#   此處只改「遇到無法編碼之字元時的行為」(改為以 ? 取代),不動編碼本身,
#   故主控台顯示維持正常。
import sys as _sys
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(errors="replace")
    except Exception:
        pass
# ---------------------------------------------------------------------------
import argparse
import csv
import decimal
import glob
import json
import os
import re
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
D = decimal.Decimal


# ======================================================================
# 〔稽核修正 DOCS-17〕目錄搜尋:以標記檔辨識,不寫死資料夾名稱
# ======================================================================
_SKIP = {"__pycache__", ".git", "Shanghai_T2DM", "figures_v2", "reference", "reference_teacher",
         "eda_figures", "cgm_output", "output", "output_reg", "output_risk", "output_cost_mc"}


def _is_deliv(d):
    return bool(glob.glob(os.path.join(d, "專題書面報告_v*.docx")))


def _is_out(d):
    return os.path.isfile(os.path.join(d, "output_risk", "quadrant_summary.csv"))


def _find(pred, explicit=""):
    if explicit:
        p = os.path.abspath(explicit)
        return p if pred(p) else ""
    d = HERE
    for _ in range(5):
        frontier, budget = [(d, 0)], 4000
        while frontier:
            x, dep = frontier.pop(0)
            budget -= 1
            if budget < 0:
                break
            if pred(x):
                return os.path.abspath(x)
            if dep >= 4:
                continue
            try:
                subs = sorted(e.path for e in os.scandir(x) if e.is_dir() and e.name not in _SKIP
                              and not e.name.startswith("reference_old_"))
            except OSError:
                continue
            frontier.extend((s, dep + 1) for s in subs)
        parent = os.path.dirname(d)
        if parent == d:
            break
        d = parent
    return ""


# [修正 31] 交付物檔名帶版本號,取版本號最大者並印出實際選用之檔名。
def pick(base, pattern):
    cs = sorted(glob.glob(os.path.join(base, pattern)))
    if not cs:
        return ""

    def ver(p):
        m = re.search(r"_v(\d+)", os.path.basename(p))
        return int(m.group(1)) if m else 0
    return max(cs, key=ver)


def deliverables(base):
    d = {"報告": pick(base, "專題書面報告_v*.docx"),
         "簡報": pick(base, "專題簡報_v*.pptx"),
         "海報": pick(base, "專題海報_80x100cm_v*.pptx")}
    appdir = pick(base, "糖三臟決策支援App_v*")
    d["App"] = os.path.join(appdir, "index.html") if appdir else ""
    d["App2"] = os.path.join(appdir, "糖三臟決策支援App.html") if appdir else ""
    return d


# ======================================================================
# 文字正規化:只換「一對一」之字元,使正規化後之索引與原文相同(前後文取原文)
# ======================================================================
_TRANS = str.maketrans({"−": "-", "－": "-", "％": "%", "．": ".",
                        **{chr(0xFF10 + i): str(i) for i in range(10)}})


def norm(t):
    return (t or "").translate(_TRANS)


# ======================================================================
# 文字單位
# ======================================================================
class Unit:
    __slots__ = ("dv", "loc", "raw", "text", "ctx", "prv", "nxt", "kind")

    def __init__(self, dv, loc, raw, kind, ctx="", prv="", nxt=""):
        self.dv, self.loc, self.raw, self.kind = dv, loc, raw, kind
        self.text = norm(raw)
        self.ctx, self.prv, self.nxt = norm(ctx), norm(prv), norm(nxt)


def _cell_texts(row):
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(c.text.strip())
    return out


def units_docx(path):
    import docx
    from docx.oxml.ns import qn
    d = docx.Document(path)
    U = []
    for i, p in enumerate(d.paragraphs):
        if p.text.strip():
            U.append(Unit("報告", f"¶{i}", p.text, "段"))
    body = list(d.element.body.iterchildren())
    caps, ti = [], 0
    for i, ch in enumerate(body):
        if ch.tag != qn("w:tbl"):
            continue
        cap = ""
        for j in range(i - 1, max(i - 4, -1), -1):
            if body[j].tag == qn("w:p"):
                t = "".join(x.text or "" for x in body[j].iter(qn("w:t"))).strip()
                if t.startswith("表 "):
                    cap = t
                    break
                if t:
                    break
        caps.append(cap)
        ti += 1
    for ti, t in enumerate(d.tables):
        cap = caps[ti] if ti < len(caps) else ""
        tag = (cap.split("　")[0].strip() if cap else "表") + f"(T{ti})"
        rows = [_cell_texts(r) for r in t.rows]
        hdr = rows[0] if rows else []
        for ri, cells in enumerate(rows):
            for ci, tx in enumerate(cells):
                if not tx:
                    continue
                rowlab = " / ".join(cells[:min(ci, 2)])
                colh = hdr[ci] if ci < len(hdr) else ""
                U.append(Unit("報告", f"{tag} r{ri}c{ci}", tx, "表", ctx=f"{cap}｜{rowlab}｜{colh}"))
    for k, tx in enumerate(d.element.body.iter(qn("w:txbxContent"))):
        t = "".join(x.text or "" for x in tx.iter(qn("w:t")))
        if t.strip():
            U.append(Unit("報告", f"文字框{k}", t, "框"))
    return U


def units_pptx(path, dv):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def flat(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from flat(sh.shapes)
            else:
                yield sh
    pr = Presentation(path)
    U = []
    for si, s in enumerate(pr.slides, 1):
        shs = list(flat(s.shapes))
        texts = [sh.text_frame.text if sh.has_text_frame else "" for sh in shs]
        where = f"第{si}張 " if dv == "簡報" else ""
        for k, sh in enumerate(shs):
            if sh.has_text_frame and texts[k].strip():
                prv = next((texts[j] for j in range(k - 1, -1, -1) if texts[j].strip()), "")
                nxt = next((texts[j] for j in range(k + 1, len(shs)) if texts[j].strip()), "")
                U.append(Unit(dv, f"{where}{'本文' if dv == '簡報' else '文字框'}〈{sh.name}〉",
                              texts[k], "框", prv=prv, nxt=nxt))
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                hdr = rows[0] if rows else []
                for ri, cells in enumerate(rows):
                    for ci, tx in enumerate(cells):
                        if tx:
                            rowlab = " / ".join(cells[:min(ci, 2)])
                            colh = hdr[ci] if ci < len(hdr) else ""
                            U.append(Unit(dv, f"{where}表 r{ri}c{ci}", tx, "表",
                                          ctx=f"{dv}表｜{rowlab}｜{colh}"))
        if s.has_notes_slide:
            t = s.notes_slide.notes_text_frame.text
            if t.strip():
                U.append(Unit(dv, f"{where}備註", t, "注"))
    return U


def units_html(path, dv="App"):
    s = open(path, encoding="utf-8", errors="replace").read()
    lines = s.split("\n")
    U = []
    for li, line in enumerate(lines, 1):
        l2 = re.sub(r"data:[a-zA-Z0-9.+/-]+;base64,[A-Za-z0-9+/=]+", "data:…", line)
        l2 = re.sub(r"[A-Za-z0-9+/]{120,}={0,2}", "<base64>", l2)
        if l2.strip():
            prv = "\n".join(lines[max(0, li - 5):li - 1])
            U.append(Unit(dv, f"L{li}", l2, "行", prv=prv))
    return U


# ======================================================================
# ROUND_HALF_UP 與進位邊界偵測
# ======================================================================
def half_up(v, nd, thou=False):
    """回傳 (字串, 邊界候選或 None)。
    若 v 本身恰有 nd+1 位小數且末位為 5(即來源 CSV 已捨入,且恰在進位邊界),
    則真值可能落在邊界兩側,回傳兩個候選值。"""
    d = D(repr(float(v)))
    q = D(1).scaleb(-nd)
    s = d.quantize(q, rounding=decimal.ROUND_HALF_UP)
    b = None
    t = repr(float(v))
    if "e" not in t and "." in t:
        dec = t.split(".")[1]
        if len(dec) == nd + 1 and dec.endswith("5"):
            b = (d.quantize(q, rounding=decimal.ROUND_DOWN), s)

    def f(x):
        if thou:
            return f"{int(x):,}" if nd == 0 else f"{x:,.{nd}f}"
        return f"{x:.{nd}f}" if nd else str(int(x))
    return f(s), (tuple(f(x) for x in b) if b else None)


# ======================================================================
# 〔稽核修正 B-03〕由輸出推導期望值
#   V[key] = (值, 小數位數 nd, 千分位, 來源說明);nd=None 表示只供 App 常數比對(不格式化)
# ======================================================================
NAMES = {}
DERIVE_ERRORS = []

# 〔稽核修正 B-10〕共線性:報告所引之兩兩相關(¶521 五組＋體重/BMI、BMI/身高;¶574 尿酸與腎功能)
COLLI_PAIRS = [("eGFR", "Cr"), ("HbA1c", "GA"), ("TC", "LDL"), ("FPG", "PPG"), ("Ins_f", "Ins_2h"),
               ("Weight", "BMI"), ("BMI", "Height"), ("UA", "Cr"), ("UA", "eGFR")]
PAIR_FEATS = {"eGFR", "Cr", "HbA1c", "GA", "TC", "LDL", "FPG", "PPG", "Ins_f", "Ins_2h"}   # ¶521「前述各組」
ABL_KEYS = {frozenset(): "abl_all", frozenset({"BMI"}): "abl_noBMI", frozenset({"Height", "Weight"}): "abl_noHW",
            frozenset({"Height", "Weight", "eGFR"}): "abl_noHWe",
            frozenset({"BMI", "Height", "Weight"}): "abl_cumBHW",
            frozenset({"BMI", "Height", "Weight", "eGFR"}): "abl_cumBHWe"}
# 選用樣式:對應「預期改寫後之措辭」或「預期將刪除之錯誤敘述」,找不到出現處不視為錯誤
OPTIONAL_KEYS = {"vif1_maxpairs", "vif1_Cr", "vif1_eGFR",
                 # 〔稽核修正〕只出現在稽核修訂後之措辭(報告 ¶833、簡報第 22 張、App v3.1 之 phNote);
                 #   學生若因新數字改寫該句,樣式對不上時不應視為「找不到」而擋下驗證
                 "alr_gpt30_diff", "alr_gpt60_diff", "pool_better_per_n", "ph30_rmse_med"}
# 共線性各鍵之名稱(輸出檔不存在時「無法檢查」列仍可顯示中文名稱)
NAMES.update({f"vif0_{f}": f"VIF {f}" for f in ("Weight", "BMI", "Height")})
NAMES.update({f"vifr2_{f}": f"以其餘特徵迴歸 {f} 之 R²" for f in ("Weight", "BMI", "Height")})
NAMES.update({f"boot_min_{f}": f"拔靴少數方向次數 {f}" for f in ("Weight", "BMI", "Height")})
NAMES.update({"vif1_maxpairs": "¶521 五組高相關變數之最大 VIF", "vif1_Cr": "VIF Cr", "vif1_eGFR": "VIF eGFR",
              "vif_n_lt5": "VIF<5 之特徵數", "bmi_def_r": "BMI 與 體重÷身高² 之相關", "bmi_def_dev": "BMI 定義式最大偏差",
              "boot_n": "拔靴次數", "boot_med_lt5": "VIF<5 特徵之少數方向次數中位數",
              "abl_n": "剔除特徵比較之重複切分組數", "abl_all": "剔除特徵之 AUC 中位數(全部 24 特徵)",
              "abl_noBMI": "剔除特徵之 AUC 中位數(剔除 BMI)", "abl_noHW": "剔除特徵之 AUC 中位數(剔除身高、體重)",
              "abl_noHWe": "剔除特徵之 AUC 中位數(剔除身高、體重、eGFR)"})
# 「無法檢查」時提示所需之輸出檔(依鍵之前綴)
NEED_FILE = [(r"^(?:vif|bmi_def)", "output/collinearity_vif.csv"),
             (r"^boot_", "output/collinearity_bootstrap_signs.csv"),
             (r"^abl_", "output/collinearity_ablation_auc.csv"),
             ("^(?:" + "|".join(f"r_{a}_{b}" for a, b in COLLI_PAIRS) + ")$", "output/eda_figures/correlation_matrix.csv")]


def derive(out_dir):
    V = {}
    DERIVE_ERRORS.clear()

    def put(key, v, nd, src, name, thou=False):
        NAMES.setdefault(key, name)
        if v is None or (isinstance(v, float) and not np.isfinite(v)):
            return
        V[key] = (float(v), nd, thou, src)

    def rd(rel):
        p = os.path.join(out_dir, rel)
        return pd.read_csv(p) if os.path.exists(p) else None

    # ---- 分類模型 AUC(model_results.csv)
    try:
        M = rd("output/model_results.csv")
        if M is not None:
            aucs = {}
            for tgt, t in [("Any_Complication", "any"), ("Microvascular", "micro"), ("Macrovascular", "macro")]:
                for mdl, m in [("Logistic", "lr"), ("XGB", "xgb")]:
                    r = M[(M.Target == tgt) & M.Model.astype(str).str.startswith(mdl)]
                    if not len(r):
                        continue
                    r = r.iloc[0]
                    zh = {"any": "任一", "micro": "小血管", "macro": "大血管"}[t] + ("LR" if m == "lr" else "XGB")
                    src = f"model_results.csv::AUC_rep_*[{tgt},{mdl}]"
                    put(f"auc_{t}_{m}", r.AUC_rep_median, 3, src, f"{zh} AUC 中位數")
                    put(f"ci_{t}_{m}_lo", r.AUC_rep_lo95, 3, src, f"{zh} AUC 區間下界")
                    put(f"ci_{t}_{m}_hi", r.AUC_rep_hi95, 3, src, f"{zh} AUC 區間上界")
                    aucs[(t, m)] = r
            if aucs:
                allv = [r.AUC_rep_median for r in aucs.values()]
                lr = [r for (t, m), r in aucs.items() if m == "lr"]
                put("auc_min_all", min(allv), 3, "model_results.csv::AUC_rep_median 六者最小", "AUC 範圍下限(兩模型三目標)")
                put("auc_max_all", max(allv), 3, "model_results.csv::AUC_rep_median 六者最大", "AUC 範圍上限(兩模型三目標)")
                put("auc_min_lr", min(r.AUC_rep_median for r in lr), 3, "LR 三目標最小", "LR AUC 範圍下限")
                put("auc_max_lr", max(r.AUC_rep_median for r in lr), 3, "LR 三目標最大", "LR AUC 範圍上限")
                put("auc_min_lr2", min(r.AUC_rep_median for r in lr), 2, "LR 三目標最小(兩位)", "LR AUC 範圍下限(兩位)")
                put("auc_max_lr2", max(r.AUC_rep_median for r in lr), 2, "LR 三目標最大(兩位)", "LR AUC 範圍上限(兩位)")
                put("ci_lr_min2", min(r.AUC_rep_lo95 for r in lr), 2, "LR 三目標 lo95 最小", "LR AUC 95% 區間最低(兩位)")
                put("ci_lr_max2", max(r.AUC_rep_hi95 for r in lr), 2, "LR 三目標 hi95 最大", "LR AUC 95% 區間最高(兩位)")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('分類模型 AUC(model_results.csv)', repr(_e)[:120]))
    # ---- 迴歸
    try:
        G = rd("output_reg/reg_results.csv")
        if G is not None:
            for m, k in [("Linear", "lin"), ("XGBoost", "xgb")]:
                r = G[G.Model.astype(str).str.startswith(m)]
                if len(r):
                    put(f"r2_{k}", r.iloc[0].R2_median, 3, f"reg_results.csv::R2_median[{m}]", f"迴歸 R²({m})")
                    # 〔稽核修正 B-03〕¶599 之區間、MAE、MAPE 原無樣式涵蓋
                    for col, nm, nd in (("R2_lo95", "R² 區間下界", 3), ("R2_hi95", "R² 區間上界", 3),
                                        ("MAE_median", "MAE 中位數", 2), ("MAPE_median", "MAPE 中位數", 1)):
                        if col in r.columns:
                            put(f"{col.lower()}_{k}", r.iloc[0][col], nd, f"reg_results.csv::{col}[{m}]", f"迴歸 {nm}({m})")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('迴歸', repr(_e)[:120]))
    # ---- LSTM
    try:
        A = rd("cgm_output/lstm_vs_baseline_all.csv")
        P = rd("cgm_output/lstm_pooled_vs_baseline.csv")
        RT = rd("output_risk/patient_risk_table.csv")
        if A is not None:
            s = "cgm_output/lstm_vs_baseline_all.csv"
            put("per_skill_med", A.skill_pct.median(), 1, f"{s}::skill_pct 中位數", "逐人 LSTM 改善中位數(15 分)")
            put("per_skill_mean", A.skill_pct.mean(), 1, f"{s}::skill_pct 平均", "逐人 LSTM 改善平均")
            put("per_beat_pct", (A.lstm_rmse < A.baseline_rmse).mean() * 100, 1, f"{s}::lstm_rmse<baseline_rmse 占比", "逐人 LSTM 勝率")
            put("per_rmse_mean", A.lstm_rmse.mean(), 2, f"{s}::lstm_rmse 平均", "逐人 LSTM RMSE 平均")
            put("per_rmse_med", A.lstm_rmse.median(), 2, f"{s}::lstm_rmse 中位數", "逐人 LSTM RMSE 中位數")
            # 〔稽核修正 C-02〕¶650 示範紀錄(病患 2000)之逐人模型數字,重跑後會變
            r0 = A[A.record.astype(str) == "2000_0_20201230"]
            if len(r0):
                r0 = r0.iloc[0]
                put("p2000_rmse", r0.lstm_rmse, 2, f"{s}::lstm_rmse[2000_0_20201230]", "示範紀錄 2000 LSTM RMSE")
                if "lstm_mae" in A.columns:
                    put("p2000_mae", r0.lstm_mae, 2, f"{s}::lstm_mae[2000_0_20201230]", "示範紀錄 2000 LSTM MAE")
                put("p2000_base", r0.baseline_rmse, 2, f"{s}::baseline_rmse[2000_0_20201230]", "示範紀錄 2000 基準 RMSE")
                put("p2000_skill", r0.skill_pct, 0, f"{s}::skill_pct[2000_0_20201230]", "示範紀錄 2000 相對改善")
        if P is not None:
            s = "cgm_output/lstm_pooled_vs_baseline.csv"
            put("pool_skill_med", P.skill_pct.median(), 1, f"{s}::skill_pct 中位數", "合併 LSTM 改善中位數(15 分)")
            put("pool_beat_pct", (P.lstm_rmse < P.baseline_rmse).mean() * 100, 1, f"{s}::lstm_rmse<baseline_rmse 占比", "合併 LSTM 勝率")
            put("pool_nbeat", int((P.lstm_rmse < P.baseline_rmse).sum()), 0, f"{s}::lstm_rmse<baseline_rmse 筆數", "合併 LSTM 勝過基準筆數")
            put("pool_n", len(P), 0, f"{s}::筆數", "LSTM 監測筆數")
            put("pool_nlose", int(len(P) - (P.lstm_rmse < P.baseline_rmse).sum()), 0, f"{s}::未勝過基準筆數", "合併 LSTM 不如基準筆數")
            put("pool_rmse_mean", P.lstm_rmse.mean(), 2, f"{s}::lstm_rmse 平均", "合併 LSTM RMSE 平均")
            put("pool_rmse_med", P.lstm_rmse.median(), 2, f"{s}::lstm_rmse 中位數", "合併 LSTM RMSE 中位數")
            put("pool_skill_min", P.skill_pct.min(), 1, f"{s}::skill_pct 最小值", "合併 LSTM 改善最小值")   # 〔稽核修正 A-27〕¶667
            put("base_rmse_mean", P.baseline_rmse.mean(), 2, f"{s}::baseline_rmse 平均", "持續性基準 RMSE 平均(15 分)")
            put("base_rmse_med", P.baseline_rmse.median(), 2, f"{s}::baseline_rmse 中位數", "持續性基準 RMSE 中位數(15 分)")
            if "n_train_windows" in P.columns:
                put("pool_train_windows", int(P.n_train_windows.sum()), 0, f"{s}::n_train_windows 總和", "合併訓練視窗數", thou=True)
        if A is not None and P is not None:
            MM = A.merge(P[["record", "lstm_rmse", "skill_pct"] + (["n_points"] if "n_points" in P.columns else [])]
                         .rename(columns={"lstm_rmse": "p_rmse", "skill_pct": "p_skill"}), on="record")
            put("pool_better_per_pct", (MM.p_rmse < MM.lstm_rmse).mean() * 100, 1,
                "兩 CSV 依 record 合併::合併 lstm_rmse < 逐人 lstm_rmse 之占比", "合併優於逐人之比例")
            # 〔稽核修正 A-27〕簡報第 22 張「約九成之點（98/109）」之筆數,重跑後會變
            put("pool_better_per_n", int((MM.p_rmse < MM.lstm_rmse).sum()), 0,
                "兩 CSV 依 record 合併::合併 lstm_rmse < 逐人 lstm_rmse 之筆數", "合併優於逐人之筆數")
            # 〔稽核修正 A-27〕報告「改善中位數逾倍(a% → b%)」之倍數主張,重跑後可能不再成立
            if A.skill_pct.median() > 0:
                put("pool_per_skill_ratio", P.skill_pct.median() / A.skill_pct.median(), 2,
                    "合併 skill_pct 中位數 ÷ 逐人 skill_pct 中位數", "合併對逐人改善中位數之倍數")
            if "n_points" in MM.columns:
                put("r_len_per", np.corrcoef(MM.skill_pct, MM.n_points)[0, 1], 2,
                    "逐人 skill_pct 與 n_points 之 Pearson r", "逐人改善與監測長度之相關")
                put("r_len_pool", np.corrcoef(MM.p_skill, MM.n_points)[0, 1], 2,
                    "合併 skill_pct 與 n_points 之 Pearson r", "合併改善與監測長度之相關")
            if RT is not None and "has_complication" in RT.columns:
                X = MM.merge(RT[["record", "has_complication"]], on="record")
                for c, lab in [(0, "無"), (1, "有")]:
                    g = X[X.has_complication == c]
                    put(f"per_skill_comp{c}", g.skill_pct.median(), 1, "逐人 skill_pct 中位數(依 has_complication)", f"{lab}併發症 逐人改善中位數")
                    put(f"pool_skill_comp{c}", g.p_skill.median(), 1, "合併 skill_pct 中位數(依 has_complication)", f"{lab}併發症 合併改善中位數")
                    put(f"per_beat_comp{c}", (g.lstm_rmse < g.baseline_rmse).mean() * 100, 1, "逐人勝率(依 has_complication)", f"{lab}併發症 逐人勝率")
                    put(f"pool_beat_comp{c}", (g.p_rmse < g.baseline_rmse).mean() * 100, 1, "合併勝率(依 has_complication)", f"{lab}併發症 合併勝率")
                    put(f"n_comp{c}", len(g), 0, "依 has_complication 之筆數", f"{lab}併發症筆數")
                    put(f"pool_nbeat_comp{c}", int((g.p_rmse < g.baseline_rmse).sum()), 0, "合併勝過基準筆數(依 has_complication)", f"{lab}併發症 合併勝過筆數")
                g0, g1 = X[X.has_complication == 0], X[X.has_complication == 1]
                put("gap_per", g0.skill_pct.median() - g1.skill_pct.median(), 1, "逐人:無併發症中位數 − 有併發症中位數", "族群落差(逐人)")
                put("gap_pool", g0.p_skill.median() - g1.p_skill.median(), 1, "合併:無併發症中位數 − 有併發症中位數", "族群落差(合併)")
        # 30、60 分鐘:優先以逐筆檔全精度計算,避免時界彙總檔(已捨入至 3 位)之二次進位
        H = rd("cgm_output/lstm_pooled_by_horizon.csv")
        for h in (30, 60):
            X = rd(f"cgm_output/lstm_pooled_PH{h}m.csv")
            s = f"cgm_output/lstm_pooled_PH{h}m.csv"
            if X is not None:
                put(f"ph{h}_rmse_mean", X.lstm_rmse.mean(), 2, f"{s}::lstm_rmse 平均", f"PH{h} LSTM RMSE 平均")
                put(f"ph{h}_rmse_med", X.lstm_rmse.median(), 2, f"{s}::lstm_rmse 中位數", f"PH{h} LSTM RMSE 中位數")
                put(f"ph{h}_mae_mean", X.lstm_mae.mean(), 2, f"{s}::lstm_mae 平均", f"PH{h} LSTM MAE 平均")
                put(f"ph{h}_skill_med", X.skill_pct.median(), 1, f"{s}::skill_pct 中位數", f"PH{h} 改善中位數")
                put(f"ph{h}_beat_pct", (X.lstm_rmse < X.baseline_rmse).mean() * 100, 1, f"{s}::勝率", f"PH{h} 勝率")
                put(f"ph{h}_base_rmse_mean", X.baseline_rmse.mean(), 2, f"{s}::baseline_rmse 平均", f"PH{h} 基準 RMSE 平均")
            if H is not None:
                r = H[H.PH_min == h]
                if len(r):
                    r = r.iloc[0]
                    s = "cgm_output/lstm_pooled_by_horizon.csv"
                    if "pooled_rmse" in H.columns:
                        put(f"ph{h}_pooled_rmse", r.pooled_rmse, 2, f"{s}::pooled_rmse[{h}]", f"PH{h} 全視窗彙總 RMSE")
                    if "r2" in H.columns:
                        put(f"ph{h}_r2", r.r2, 3, f"{s}::r2[{h}]", f"PH{h} 決定係數 R²")
                    if "n_test_windows" in H.columns:
                        put(f"ph{h}_nwin", int(r.n_test_windows), 0, f"{s}::n_test_windows[{h}]", f"PH{h} 測試視窗數", thou=True)
        # 〔稽核修正 A-10〕¶833 與 Alredaini 等(2026)之比較句:以本研究彙總 RMSE(報告所列之兩位小數)與該文數值推導,
        #   重跑後「分別低 x% 與 y%」「高出 a 與 b mg/dL」及「皆略優於本研究」之主張須隨之核對
        LIT = {"lstm30": 13.38, "lstm60": 23.94, "gpt30": 12.57, "gpt60": 22.45}     # 該文表 9、表 11(文獻值)
        if "ph30_pooled_rmse" in V and "ph60_pooled_rmse" in V:
            p30 = float(half_up(V["ph30_pooled_rmse"][0], 2)[0])
            p60 = float(half_up(V["ph60_pooled_rmse"][0], 2)[0])
            s = "lstm_pooled_by_horizon.csv::pooled_rmse(兩位小數)與 Alredaini 等(2026)"
            put("alr_lstm30_rel", (LIT["lstm30"] - p30) / LIT["lstm30"] * 100, 1, s, "較該文 LSTM 低之比例(30 分,%)")
            put("alr_lstm60_rel", (LIT["lstm60"] - p60) / LIT["lstm60"] * 100, 1, s, "較該文 LSTM 低之比例(60 分,%)")
            put("alr_gpt30_diff", p30 - LIT["gpt30"], 2, s, "本研究高出 GPT-4.1 之 RMSE(30 分)")
            put("alr_gpt60_diff", p60 - LIT["gpt60"], 2, s, "本研究高出 GPT-4.1 之 RMSE(60 分)")
            put("alr_gpt_min_diff", min(p30 - LIT["gpt30"], p60 - LIT["gpt60"]), 2, s, "GPT-4.1 於兩時界皆較佳(較小之差)")
        LG = rd("output/lstm_by_length_group.csv")
        if LG is not None and len(LG) >= 4:
            s = "output/lstm_by_length_group.csv"
            for i in range(4):
                r = LG.iloc[i]
                put(f"lg{i+1}_per_skill", r.per_skill_median, 1, f"{s}::per_skill_median[{i}]", f"長度第 {i+1} 組 逐人改善")
                put(f"lg{i+1}_pool_skill", r.pooled_skill_median, 1, f"{s}::pooled_skill_median[{i}]", f"長度第 {i+1} 組 合併改善")
                put(f"lg{i+1}_per_beat", r.per_beat_pct, 1, f"{s}::per_beat_pct[{i}]", f"長度第 {i+1} 組 逐人勝率")
                put(f"lg{i+1}_pool_beat", r.pooled_beat_pct, 1, f"{s}::pooled_beat_pct[{i}]", f"長度第 {i+1} 組 合併勝率")
                put(f"lg{i+1}_mean_points", r.mean_points, 0, f"{s}::mean_points[{i}]", f"長度第 {i+1} 組 平均點數", thou=True)
            put("lg_pool_skill_min", LG.pooled_skill_median.iloc[:4].min(), 1, f"{s}::pooled_skill_median 最小", "長度分組合併改善(最小)")
            put("lg_pool_skill_max", LG.pooled_skill_median.iloc[:4].max(), 1, f"{s}::pooled_skill_median 最大", "長度分組合併改善(最大)")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('LSTM', repr(_e)[:120]))
    # ---- 診斷值(僅取 computed 類;literature/code_constant 不是本研究之計算結果)
    try:
        DG = rd("output/diagnostics_v2.csv")
        if DG is not None:
            dd = dict(zip(DG["item"], DG["value"]))
            s = "output/diagnostics_v2.csv"
            for key, item, nd, name, scale in [
                    ("insample_med_pct", "insample_median", 1, "全資料配適預測中位數(%)", 100),
                    ("flagged_old_pct", "flagged_pct_old_cut", 1, "原門檻套用 OOF 之高風險比例(%)", 1),
                    ("gini", "cost_gini", 3, "Gini 係數", 1), ("top10", "cost_top10pct_share", 1, "最貴前 10% 占比", 1),
                    ("top20", "cost_top20pct_share", 1, "最貴前 20% 占比", 1), ("top30", "cost_top30pct_share", 1, "最貴前 30% 占比", 1),
                    ("top50", "cost_top50pct_share", 1, "最貴前 50% 占比", 1),
                    ("cost_p75_wan", "cost_p75", 1, "成本第 75 百分位(萬)", 1e-4), ("cost_p90_wan", "cost_p90", 1, "成本第 90 百分位(萬)", 1e-4),
                    ("cost_max_wan", "cost_max", 1, "成本最大值(萬)", 1e-4),
                    ("corr_zero", "corr_comp_skill_zero_order", 2, "併發症與逐人改善之零階相關", 1),
                    ("partial_r", "partial_comp_skill_ctrl_npoints", 2, "偏相關(控制監測點數)", 1),
                    ("partial_log", "partial_comp_skill_ctrl_log_npoints", 2, "偏相關(控制 log 點數)", 1),
                    ("partial_rank", "partial_comp_skill_rank", 2, "等級偏相關", 1),
                    ("corr_hba1c", "corr_npoints_hba1c", 3, "監測點數與 HbA1c 相關", 1),
                    ("npoints_comp1", "npoints_mean_with_comp", 0, "有併發症者平均監測點數", 1),
                    ("npoints_comp0", "npoints_mean_without_comp", 0, "無併發症者平均監測點數", 1),
                    ("cpi_cum", "cpi_cum_2024", 4, "CPI 累計倍數(2001→2024)", 1)]:
                if item in dd:
                    put(key, dd[item] * scale, nd, f"{s}::{item}", name, thou=key.startswith("npoints"))
            for k in ("none", "macro", "micro", "both"):
                if f"cost_direct_{k}" in dd:
                    put(f"cost_direct_{k}", dd[f"cost_direct_{k}"], None, f"{s}::cost_direct_{k}(程式常數)", f"年直接醫療費用({k})")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('診斷值(僅取 computed 類;literature/code_constant 不是本研究之計算結果)', repr(_e)[:120]))
    # ---- 四象限(紀錄)
    try:
        Q = rd("output_risk/quadrant_summary.csv")
        if Q is not None:
            for _, x in Q.iterrows():
                k = str(x["quadrant"])[0]
                s = "output_risk/quadrant_summary.csv"
                put(f"qn_{k}", x["人數"], 0, f"{s}::人數[{k}]", f"象限 {k} 人數(紀錄)")
                put(f"qcomp_{k}", x["實際併發症比例"], 1, f"{s}::實際併發症比例[{k}]", f"象限 {k} 併發症比例")
                put(f"qlong_{k}", x["平均長期風險"], 1, f"{s}::平均長期風險[{k}]", f"象限 {k} 平均長期風險")
                put(f"qshort_{k}", x["平均短期風險"], 1, f"{s}::平均短期風險[{k}]", f"象限 {k} 平均短期風險")
        if RT is not None:
            s = "output_risk/patient_risk_table.csv"
            put("cut_long_pct", RT.long_risk.median() * 100, 1, f"{s}::long_risk 中位數×100", "長期軸切點(%)")
            put("cut_long_raw", RT.long_risk.median(), None, f"{s}::long_risk 中位數", "長期軸切點")
            put("cut_short_pct", RT.short_risk.median(), 2, f"{s}::short_risk 中位數", "短期軸切點(%)")
            put("cut_short_raw", RT.short_risk.median(), None, f"{s}::short_risk 中位數", "短期軸切點")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('四象限(紀錄)', repr(_e)[:120]))
    # ---- 成本
    try:
        QC = rd("output_cost_mc/quadrant_cost_interval.csv")
        if QC is not None:
            s = "output_cost_mc/quadrant_cost_interval.csv"
            for _, x in QC.iterrows():
                k = str(x["quadrant"])[0]
                put(f"qcn_{k}", x["n"], 0, f"{s}::n[{k}]", f"象限 {k} 人數(病患)")
                put(f"qcost_{k}", x["mean"] / 1e4, 1, f"{s}::mean[{k}]÷10⁴", f"象限 {k} 平均成本(萬)")
                put(f"qcost_raw_{k}", x["mean"], None, f"{s}::mean[{k}]", f"象限 {k} 平均成本(元)")
                put(f"qlo_{k}", x["lo"] / 1e4, 1, f"{s}::lo[{k}]÷10⁴", f"象限 {k} 區間下界(萬)")
                put(f"qhi_{k}", x["hi"] / 1e4, 1, f"{s}::hi[{k}]÷10⁴", f"象限 {k} 區間上界(萬)")
        C2 = rd("output_cost_mc/patient_lifetime_cost_mc.csv")
        if C2 is not None and "quadrant" in C2.columns:
            f = C2[C2.counted_in_population.astype(bool)]
            q = f.quadrant.astype(str).str[0]
            hi, lo = f.lifetime_mean[q.isin(["A", "B"])].mean(), f.lifetime_mean[q.isin(["C", "D"])].mean()
            put("cost_ratio_ABCD", hi / lo, 2,
                "patient_lifetime_cost_mc.csv::lifetime_mean(首筆 100 位),A、B 之每人平均 ÷ C、D 之每人平均",
                "(A+B)/(C+D) 成本倍數")
        S = rd("output_cost_mc/cost_sensitivity_class_weight.csv")
        if S is not None:
            s = "output_cost_mc/cost_sensitivity_class_weight.csv"
            b = S[S.class_weight.astype(str) == "balanced"]
            nw = S[S.class_weight.isna()]
            if len(b):
                b = b.iloc[0]
                put("cost_mean", b.cost_mean_first_only, 0, f"{s}::cost_mean_first_only", "每人平均餘生成本(元)", thou=True)
                put("cost_median", b.cost_median_first_only, 0, f"{s}::cost_median_first_only", "餘生成本中位數(元)", thou=True)
                put("cost_median_wan", b.cost_median_first_only / 1e4, 1, f"{s}::cost_median_first_only÷10⁴", "餘生成本中位數(萬)")
                put("cost_mean_wan", b.cost_mean_first_only / 1e4, 1, f"{s}::cost_mean_first_only÷10⁴", "每人平均餘生成本(萬)")
                put("cost_mean_rec", b.cost_mean_per_record, 0, f"{s}::cost_mean_per_record", "每筆紀錄平均成本(元)", thou=True)
                put("pop_total_wan", b.pop_total / 1e4, 0, f"{s}::pop_total÷10⁴", "母體總成本(萬)", thou=True)
                put("pop_lo_wan", b.pop_lo / 1e4, 0, f"{s}::pop_lo÷10⁴", "母體總成本下界(萬)", thou=True)
                put("pop_hi_wan", b.pop_hi / 1e4, 0, f"{s}::pop_hi÷10⁴", "母體總成本上界(萬)", thou=True)
            if len(b):
                put("macro_auc_bal", b.macro_auc, 3, f"{s}::macro_auc[balanced]", "大血管 AUC(balanced,成本用)")
                put("micro_auc_bal", b.micro_auc, 3, f"{s}::micro_auc[balanced]", "小血管 AUC(balanced,成本用)")
            if len(nw):
                nw = nw.iloc[0]
                put("macro_pred_nw", nw.macro_pred, 1, f"{s}::macro_pred[未加權]", "大血管預測均值(未加權)")
                put("micro_pred_nw", nw.micro_pred, 1, f"{s}::micro_pred[未加權]", "小血管預測均值(未加權)")
                put("macro_auc_nw", nw.macro_auc, 3, f"{s}::macro_auc[未加權]", "大血管 AUC(未加權)")
                put("micro_auc_nw", nw.micro_auc, 3, f"{s}::micro_auc[未加權]", "小血管 AUC(未加權)")
                put("cost_mean_wan_nw", nw.cost_mean_first_only / 1e4, 1, f"{s}::cost_mean_first_only[未加權]÷10⁴", "未加權情境 每人平均(萬)")
                put("pop_total_wan_nw", nw.pop_total / 1e4, 0, f"{s}::pop_total[未加權]÷10⁴", "未加權情境 母體總成本(萬)", thou=True)
            if len(b):
                put("macro_prev", b.macro_prev, 1, f"{s}::macro_prev", "大血管實際盛行率(%)")
                put("micro_prev", b.micro_prev, 1, f"{s}::micro_prev", "小血管實際盛行率(%)")
                put("macro_pred_bal", b.macro_pred, 1, f"{s}::macro_pred[balanced]", "大血管預測均值(balanced)")
                put("micro_pred_bal", b.micro_pred, 1, f"{s}::micro_pred[balanced]", "小血管預測均值(balanced)")
                put("macro_pred_gap", b.macro_pred - b.macro_prev, 1, f"{s}::macro_pred−macro_prev", "大血管預測均值高出盛行率(百分點)")
                put("micro_pred_gap", b.micro_pred - b.micro_prev, 1, f"{s}::micro_pred−micro_prev", "小血管預測均值高出盛行率(百分點)")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('成本', repr(_e)[:120]))
    # ---- 特徵集敏感度(表 5-4-1;sensitivity_missing_features.csv,LR 40 次重複)
    try:
        SM = rd("output/sensitivity_missing_features.csv")
        if SM is not None:
            s = "output/sensitivity_missing_features.csv"
            for _, x in SM.iterrows():
                if not str(x["model"]).startswith("LR"):
                    continue
                nf = int(x["n_features"])
                t = {"Any_Complication": "any", "Microvascular": "micro", "Macrovascular": "macro"}.get(x["target"], x["target"])
                for c, lab in (("AUC_median", "med"), ("AUC_lo95", "lo"), ("AUC_hi95", "hi")):
                    put(f"sens{nf}_{t}_{lab}", x[c], 3, f"{s}::{c}[{nf} 特徵,{x['target']},LR]",
                        f"特徵集 {nf} 個 {t} LR AUC {lab}")
                put(f"sens{nf}_{t}_w", x["AUC_hi95"] - x["AUC_lo95"], 3, f"{s}::AUC_hi95−AUC_lo95[{nf} 特徵,{x['target']},LR]",
                    f"特徵集 {nf} 個 {t} LR 區間寬度")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('特徵集敏感度(表 5-4-1;sensitivity_missing_features.csv,LR 40 次重複)', repr(_e)[:120]))
    # ---- 馬可夫
    try:
        MS = rd("cgm_output/markov_stationary_3state.csv")
        if MS is not None:
            r = MS.iloc[0]
            s = "cgm_output/markov_stationary_3state.csv"
            cols = list(MS.columns)
            put("st3_low", r[cols[0]], 1, f"{s}::{cols[0]}", "三態穩態 低血糖(%)")
            put("st3_in", r[cols[1]], 1, f"{s}::{cols[1]}", "三態穩態 範圍內(%)")
            put("st3_high", r[cols[2]], 1, f"{s}::{cols[2]}", "三態穩態 高血糖(%)")
            for i, c in enumerate(cols[:3]):
                put(f"st3_frac{i}", r[c] / 100.0, None, f"{s}::{c}÷100", f"三態穩態比例 {i}")
            put("lambda2_3", r.lambda2, 3, f"{s}::lambda2", "三態 |λ₂|")
            put("halflife3_h", r.half_life_hours, 1, f"{s}::half_life_hours", "三態半衰期(小時)")
        MS5 = rd("cgm_output/markov_stationary_5state.csv")
        if MS5 is not None:
            r = MS5.iloc[0]
            s = "cgm_output/markov_stationary_5state.csv"
            inr = [c for c in MS5.columns if c.startswith("InRange")]
            if inr:
                put("st5_in", r[inr[0]], 1, f"{s}::{inr[0]}", "五態穩態 範圍內(%)")
            put("halflife5_h", r.half_life_hours, 1, f"{s}::half_life_hours", "五態半衰期(小時)")
        M5 = rd("cgm_output/markov_5state.csv")
        if M5 is not None:
            vals = M5.iloc[:, 1:6].to_numpy(dtype=float)
            for i in range(5):
                for j in range(5):
                    put(f"m5_{i}{j}", vals[i, j], 3, f"cgm_output/markov_5state.csv::[{i},{j}]", f"五態轉移機率 {i}→{j}")
        CL = rd("output/cluster_complication_rate.csv")
        if CL is not None:
            s = "output/cluster_complication_rate.csv"
            km = CL[CL.method.astype(str).str.startswith("KMeans")]
            for tgt, k in (("Macrovascular", "macro"), ("Any_Complication", "any"), ("Microvascular", "micro")):
                for c in (0, 1):
                    r = km[(km.target == tgt) & (km.cluster == c)]
                    if len(r):
                        put(f"clu_{k}{c}", r.iloc[0].pct_positive, 1, f"{s}::pct_positive[KMeans,{tgt},{c}]", f"K-means 群 {c} {tgt} 比例")
        M3 = rd("cgm_output/markov_3state.csv")
        if M3 is not None:
            vals = M3.iloc[:, 1:4].to_numpy(dtype=float)
            for i in range(3):
                for j in range(3):
                    put(f"m3_{i}{j}", vals[i, j], 3, f"cgm_output/markov_3state.csv::[{i},{j}]", f"三態轉移機率 {i}→{j}")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('馬可夫', repr(_e)[:120]))
    # ---- CGM 族群中位數
    try:
        CG = rd("cgm_output/cgm_metrics.csv")
        if CG is not None:
            s = "cgm_output/cgm_metrics.csv"
            for c, k, nd in [("TIR", "tir", 1), ("TAR", "tar", 1), ("TBR", "tbr", 1), ("CV", "cv", 1), ("GMI", "gmi", 2)]:
                if c in CG.columns:
                    put(f"{k}_med", CG[c].median(), nd, f"{s}::{c} 中位數", f"CGM {c} 中位數")
                    for stat, fn in [("mean", CG[c].mean()), ("p25", CG[c].quantile(.25)),
                                     ("p50", CG[c].median()), ("p75", CG[c].quantile(.75))]:
                        put(f"{k}_{stat}_raw", fn, None, f"{s}::{c} {stat}", f"CGM {c} {stat}")
            if "mean" in CG.columns:
                for stat, fn in [("mean", CG["mean"].mean()), ("p25", CG["mean"].quantile(.25)),
                                 ("p50", CG["mean"].median()), ("p75", CG["mean"].quantile(.75))]:
                    put(f"glu_{stat}_raw", fn, None, f"{s}::mean {stat}", f"CGM 平均血糖 {stat}")
    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('CGM 族群中位數', repr(_e)[:120]))
    # ---- 勝算比
    try:
        for tgt, t in [("Any_Complication", "any"), ("Microvascular", "micro"), ("Macrovascular", "macro")]:
            O = rd(f"output/lr_oddsratio_{tgt}.csv")
            if O is not None:
                for _, x in O.iterrows():
                    put(f"or_{t}_{x['Feature']}", x["OddsRatio"], 2, f"lr_oddsratio_{tgt}.csv::OddsRatio[{x['Feature']}]",
                        f"勝算比 {t} {x['Feature']}")

    except Exception as _e:                      # 〔稽核修正 B-03〕單一區段失敗不影響其他區段
        DERIVE_ERRORS.append(('勝算比', repr(_e)[:120]))
    # ---- 〔稽核修正 B-10〕共線性(報告 ¶521、¶574、¶585、¶802、¶803)
    #   相關係數:correlation_matrix.csv(原始臨床單位、成對完整觀測);其餘:collinearity_check.py 之輸出。
    #   四個檔各自獨立推導,任一檔缺欄只影響該檔之項目。
    try:
        CM = rd("output/eda_figures/correlation_matrix.csv")
        if CM is not None:
            CM = CM.set_index(CM.columns[0])
            s = "output/eda_figures/correlation_matrix.csv"
            for a_, b_ in COLLI_PAIRS:
                if a_ in CM.index and b_ in CM.columns:
                    put(f"r_{a_}_{b_}", CM.loc[a_, b_], 2, f"{s}::[{a_},{b_}]", f"相關係數 {a_}–{b_}")
    except Exception as _e:
        DERIVE_ERRORS.append(('相關係數(correlation_matrix.csv)', repr(_e)[:120]))
    try:
        VF = rd("output/collinearity_vif.csv")
        if VF is not None:
            s = "output/collinearity_vif.csv"
            vv = VF[VF.section.astype(str) == "VIF"].set_index("feature")
            for f_ in ("Weight", "BMI", "Height"):
                put(f"vif0_{f_}", vv.loc[f_, "value"], 0, f"{s}::value[VIF,{f_}]", f"VIF {f_}")
                put(f"vifr2_{f_}", vv.loc[f_, "r2"], 3, f"{s}::r2[VIF,{f_}]", f"以其餘特徵迴歸 {f_} 之 R²")
            for f_ in ("Cr", "eGFR"):
                put(f"vif1_{f_}", vv.loc[f_, "value"], 1, f"{s}::value[VIF,{f_}]", f"VIF {f_}")
            pv = vv.loc[[f_ for f_ in vv.index if f_ in PAIR_FEATS], "value"].astype(float)
            if len(pv):
                fm = str(pv.idxmax())
                put("vif1_maxpairs", pv.max(), 1, f"{s}::value[VIF] 於 ¶521 所列五組變數中之最大者({fm})", "")
                NAMES["vif1_maxpairs"] = f"¶521 五組高相關變數之最大 VIF(為 {fm})"
            sm = VF[VF.section.astype(str) == "VIF_summary"].set_index("feature")
            put("vif_n_lt5", sm.loc["n_features_VIF_lt5", "value"], 0,
                f"{s}::value[VIF_summary,n_features_VIF_lt5]", "VIF<5 之特徵數")
            bd = VF[VF.section.astype(str) == "BMI_definition"].set_index("feature")
            put("bmi_def_r", bd.loc["BMI", "value"], 6, f"{s}::value[BMI_definition,BMI]", "BMI 與 體重÷身高² 之相關")
            put("bmi_def_dev", bd.loc["max_abs_dev", "value"], 3, f"{s}::value[BMI_definition,max_abs_dev]",
                "BMI 定義式最大偏差")
    except Exception as _e:
        DERIVE_ERRORS.append(('共線性 VIF(collinearity_vif.csv)', repr(_e)[:120]))
    try:
        BS = rd("output/collinearity_bootstrap_signs.csv")
        if BS is not None:
            s = "output/collinearity_bootstrap_signs.csv"
            br = BS[BS.resample_unit.astype(str) == "record"].set_index("feature")
            put("boot_n", br.n_boot.iloc[0], 0, f"{s}::n_boot", "拔靴次數")
            for f_ in ("Weight", "BMI", "Height"):
                put(f"boot_min_{f_}", br.loc[f_, "minority_count"], 0, f"{s}::minority_count[record,{f_}]",
                    f"拔靴少數方向次數 {f_}")
            put("boot_med_lt5", br.loc["median_VIF_lt5", "minority_count"], 1,
                f"{s}::minority_count[record,median_VIF_lt5]", "VIF<5 特徵之少數方向次數中位數")
    except Exception as _e:
        DERIVE_ERRORS.append(('共線性拔靴(collinearity_bootstrap_signs.csv)', repr(_e)[:120]))
    try:
        AB = rd("output/collinearity_ablation_auc.csv")
        if AB is not None:
            s = "output/collinearity_ablation_auc.csv"
            put("abl_n", AB.n_repeats.iloc[0], 0, f"{s}::n_repeats", "剔除特徵比較之重複切分組數")
            for _, x in AB.iterrows():
                rm = frozenset(t for t in re.split(r"[、,，\s]+", str(x["removed"])) if t and t != "(無)")
                if rm in ABL_KEYS:
                    put(ABL_KEYS[rm], x["auc_median"], 3, f"{s}::auc_median[{x['feature_set']}]",
                        f"剔除特徵之 AUC 中位數({x['feature_set']})")
    except Exception as _e:
        DERIVE_ERRORS.append(('剔除特徵 AUC(collinearity_ablation_auc.csv)', repr(_e)[:120]))
    return V


# ======================================================================
# 〔稽核修正 B-03〕語境樣式
#   <鍵:型別> 會展開為具名群組;同一樣式中同鍵出現兩次時以 鍵__2 區分。
#   型別:N1/N2/N3/N4/N6＝1/2/3/4/6 位小數(可帶正負號);NT＝整數(可含千分位);I＝整數;C＝「九成」類主張
# ======================================================================
FRAG = {
    "N1": r"(?<![\d.])[-+]?\d{1,4}\.\d(?![\d])",
    "N2": r"(?<![\d.])[-+]?\d{1,4}\.\d{2}(?![\d])",
    "N3": r"(?<![\d.])[-+]?\d\.\d{3}(?![\d])",
    "N4": r"(?<![\d.])[-+]?\d\.\d{4}(?![\d])",
    "N6": r"(?<![\d.])[-+]?\d\.\d{6}(?![\d])",                  # 〔稽核修正 B-10〕相關係數 1.000000
    "NT": r"(?<![\d.,])\d{1,3}(?:,\d{3})+(?![\d,])|(?<![\d.,])\d+(?![\d.,])",
    "I": r"(?<![\d.])\d+(?![\d.])",
    "C": r"逾九成|超過九成|九成以上|近九成|約九成",
    # 〔稽核修正 A-27〕「倍數」類主張;最後一項容許改寫為不含倍數之措辭(如「明顯提升」),此時只核對數字
    "R": r"逾倍|超過一倍|逾一倍|倍增|[^（）\d]{0,8}?",
    # 〔稽核修正 A-10〕「兩時界皆優於」類主張
    "G": r"皆略優於|皆優於|均略優於|均優於",
}


class Pat:
    def __init__(self, rx, dv=None, kind=None, tbl=None, prv=None, nxt=None, need=None):
        src = re.sub(r"<(\w+):(\w+)>", lambda m: f"(?P<{m.group(1)}>{FRAG[m.group(2)]})", rx)
        self.src = rx
        self.rx = re.compile(src)
        self.dv = set(dv) if dv else None
        self.kind = set(kind) if kind else None
        self.tbl = tuple(re.compile(x) for x in tbl) if tbl else None
        self.prv = re.compile(prv) if prv else None
        self.nxt = re.compile(nxt) if nxt else None
        self.need = re.compile(need) if need else None

    def applies(self, u):
        if self.dv and u.dv not in self.dv:
            return False
        if self.kind and u.kind not in self.kind:
            return False
        if self.tbl:
            parts = u.ctx.split("｜")
            if len(parts) != 3 or not all(r.search(x) for r, x in zip(self.tbl, parts)):
                return False
        if self.prv and not self.prv.search(u.prv):
            return False
        if self.nxt and not self.nxt.search(u.nxt):
            return False
        if self.need and not self.need.search(u.text):
            return False
        return True


def TBL(cap, row, col, vrx):
    """表格格位樣式:表題、該列左側格位(以「 / 」連接)、欄首三段分別比對。"""
    return Pat(vrx, kind={"表"}, tbl=(cap, row, col))


_Q = {"A": ("A", r"A　?立即|^A$|^A / ", r"長高＋短高"), "B": ("B", r"B　?慢性|^B$|^B / ", r"長高＋短低"),
      "C": ("C", r"C　?血糖|^C$|^C / ", r"長低＋短高"), "D": ("D", r"D　?常規|^D$|^D / ", r"長低＋短低")}

PATTERNS = [
    # ---------------- AUC
    Pat(r"AUC\s*介於\s*<auc_min_all:N3>\s*至\s*<auc_max_all:N3>"),
    Pat(r"AUC ranged from <auc_min_all:N3> to <auc_max_all:N3>"),
    Pat(r"其值為\s*<auc_any_lr:N3>\s*（100\s*次"),
    Pat(r"AUC\s*中位數（<auc_any_lr:N3>\s*對\s*<auc_any_xgb:N3>）"),
    Pat(r"Logistic Regression\s*較高（<auc_any_lr:N3>\s*對\s*<auc_any_xgb:N3>、<auc_macro_lr:N3>\s*對\s*<auc_macro_xgb:N3>）"),
    Pat(r"XGBoost\s*略高（<auc_micro_xgb:N3>\s*對\s*<auc_micro_lr:N3>）"),
    Pat(r"以其略高（<auc_micro_xgb:N3>\s*對\s*<auc_micro_lr:N3>）"),
    Pat(r"XGBoost略高[—–-]+<auc_micro_xgb:N3>\s*對\s*<auc_micro_lr:N3>"),
    Pat(r"對應其\s*AUC\s*<auc_macro_lr:N3>，相對於\s*XGBoost\s*之\s*<auc_macro_xgb:N3>"),
    Pat(r"^\s*AUC\s*<auc_min_lr:N3>\s*[–—-]\s*<auc_max_lr:N3>\s*$", kind={"框"}),
    Pat(r"判別效能\s*<auc_min_lr2:N2>\s*到\s*<auc_max_lr2:N2>"),
    Pat(r"任一\s*<auc_any_lr:N3>\s*［<ci_any_lr_lo:N3>,\s*<ci_any_lr_hi:N3>］、小血管\s*<auc_micro_lr:N3>\s*"
        r"［<ci_micro_lr_lo:N3>,\s*<ci_micro_lr_hi:N3>］、大血管\s*<auc_macro_lr:N3>\s*［<ci_macro_lr_lo:N3>,\s*<ci_macro_lr_hi:N3>］"),
    Pat(r"區間(?:（[^）]{0,30}）)?高度重疊（<ci_micro_xgb_lo:N3>\s*[–—-]\s*<ci_micro_xgb_hi:N3>\s*對\s*<ci_micro_lr_lo:N3>\s*[–—-]\s*<ci_micro_lr_hi:N3>）"),
    Pat(r"(?:95%\s*區間|100\s*次重複之?區間)\s*<ci_lr_min2:N2>\s*[–—-]\s*<ci_lr_max2:N2>"),   # 〔稽核修正 A-12〕接受改寫後之措辭
] + [
    TBL(r"表 4-3-1", rl, r"AUC", rf"^<{k}:N3>$")
    for rl, k in [(r"任一併發症 / Logistic", "auc_any_lr"), (r"任一併發症 / XGBoost", "auc_any_xgb"),
                  (r"小血管併發症 / Logistic", "auc_micro_lr"), (r"小血管併發症 / XGBoost", "auc_micro_xgb"),
                  (r"大血管併發症 / Logistic", "auc_macro_lr"), (r"大血管併發症 / XGBoost", "auc_macro_xgb")]
] + [
    TBL(r"簡報表", rl, cl, rf"^<{k}:N3>$")   # 〔2026-09-29 用詞統一〕表頭接受「邏輯斯迴歸 AUC」或「Logistic Regression AUC」
    for rl, cl, k in [(r"任一併發症", r"(?:邏輯斯迴歸|Logistic Regression) AUC", "auc_any_lr"), (r"任一併發症", r"XGBoost AUC", "auc_any_xgb"),
                      (r"大血管", r"(?:邏輯斯迴歸|Logistic Regression) AUC", "auc_macro_lr"), (r"大血管", r"XGBoost AUC", "auc_macro_xgb"),
                      (r"小血管", r"(?:邏輯斯迴歸|Logistic Regression) AUC", "auc_micro_lr"), (r"小血管", r"XGBoost AUC", "auc_micro_xgb")]
] + [
    # ---------------- 特徵集敏感度(表 5-4-1)
] + [
    TBL(r"表 5-4-1", rl, cl, rf"^<sens{nf}_{t}_med:N3>［<sens{nf}_{t}_lo:N3>,\s*<sens{nf}_{t}_hi:N3>］$")
    for rl, nf in [(r"缺失\s*<10%", 10), (r"缺失\s*<20%", 13), (r"缺失\s*<30%", 21), (r"全部", 24)]
    for cl, t in [(r"任一併發症", "any"), (r"小血管", "micro"), (r"大血管", "macro")]
] + [
    # ---------------- 迴歸
    Pat(r"R²\s*由\s*<r2_lin:N3>\s*［<r2_lo95_lin:N3>,\s*<r2_hi95_lin:N3>］\s*提升至\s*<r2_xgb:N3>\s*［<r2_lo95_xgb:N3>,\s*<r2_hi95_xgb:N3>］"),   # 〔稽核修正 B-03〕區間一併核對
    Pat(r"R²\s*由\s*<r2_lin:N3>\s*［[^］]*］\s*提升至\s*<r2_xgb:N3>"),
    Pat(r"MAE\s*由\s*<mae_median_lin:N2>\s*降至\s*<mae_median_xgb:N2>、MAPE\s*由\s*<mape_median_lin:N1>%\s*降至\s*<mape_median_xgb:N1>%"),
    TBL(r"表 4-6-1", r"Linear", r"R²", r"^<r2_lin:N3>$"),
    TBL(r"表 4-6-1", r"XGBoost", r"R²", r"^<r2_xgb:N3>$"),
    # ---------------- LSTM:15 分鐘(next-step)
    Pat(r"(?:改善中位數|技巧分數中位數)(?:（[^）]{0,15}）)?(?:由|自|從)?\s*<per_skill_med:N1>%?\s*(?:提升至|升至|提升到|拉到)\s*<pool_skill_med:N1>%"),   # 〔稽核修正 A-27〕容許「（預測時界 15 分鐘）」
    Pat(r"改善中位數<claim_skill_ratio:R>（<per_skill_med:N1>%\s*→\s*<pool_skill_med:N1>%）"),   # 〔稽核修正 A-27〕並核對「逾倍」
    Pat(r"from <per_skill_med:N1>% to <pool_skill_med:N1>%"),
    Pat(r"^\s*<per_skill_med:N1>%\s*→\s*<pool_skill_med:N1>%\s*$", kind={"框"}, nxt=r"訓練"),
    Pat(r"有併發症之\s*<n_comp1:I>\s*筆監測中\s*<pool_nbeat_comp1:I>\s*筆優於基準"),
    Pat(r"(?<!之)(?<!之\s)<pool_n:I>\s*筆監測中\s*<pool_nbeat:I>\s*筆(?:優於|勝過|受益)"),
    Pat(r"with <pool_nbeat:I> of <pool_n:I> recordings"),
    Pat(r"<pool_nbeat:I>\s*[/／]\s*<pool_n:I>\s*筆監測(?:受益|優於持續性基準|優於基準|勝過基準)"),   # 〔稽核修正 A-27〕改寫後之措辭
    Pat(r"勝過基準之比例自\s*<per_beat_pct:N1>%\s*升至\s*<pool_beat_pct:N1>%"),
    Pat(r"均方根誤差(?:（[^）]{0,10}）)?自\s*<per_rmse_mean:N2>\s*降至\s*<pool_rmse_mean:N2>\s*mg/dL"),   # 〔稽核修正 A-27〕容許「（逐紀錄平均）」
    Pat(r"數值（改善中位數\s*<pool_skill_med:N1>%、勝率\s*<pool_beat_pct:N1>%）"),
    Pat(r"15\s*分鐘時基準\s*RMSE\s*<base_rmse_mean:N2>、LSTM\s*<pool_rmse_mean:N2>、改善中位數\s*<pool_skill_med:N1>%、勝率\s*<pool_beat_pct:N1>%"),
    Pat(r"合併訓練之勝率分別為\s*<pool_beat_pct:N1>%\s*與"),
    Pat(r"改善中位數分別為\s*<pool_skill_med:N1>%\s*與"),
    Pat(r"逐人專屬模型之勝率分別為\s*<per_beat_pct:N1>%\s*與"),
    Pat(r"仍有\s*<pool_skill_med:N1>%\s*之改善"),
    Pat(r"RMSE\s*平均\s*<per_rmse_mean:N2>\s*mg/dL（中位數\s*<per_rmse_med:N2>）"),
    Pat(r"基準平均\s*<base_rmse_mean:N2>\s*mg/dL"),
    Pat(r"相對改善中位數\s*<per_skill_med:N1>%、平均\s*<per_skill_mean:N1>%"),
    Pat(r"於\s*<per_beat_pct:N1>%\s*之監測上優於基準"),
    Pat(r"平均改善率（<per_skill_mean:N1>%）明顯低於中位數（<per_skill_med:N1>%）"),
    Pat(r"誤差從\s*<base_rmse_mean:N2>\s*降到\s*<pool_rmse_mean:N2>"),
    Pat(r"誤差中位數從\s*<base_rmse_med:N2>\s*降到\s*<pool_rmse_med:N2>"),   # 〔稽核修正 A-27〕簡報第 22 張備註改寫後
    # 〔稽核修正 C-02〕¶650 示範紀錄;〔稽核修正 A-27〕¶667 改善最小值
    Pat(r"以病患\s*2000\s*為例，LSTM\s*測試段\s*RMSE\s*為\s*<p2000_rmse:N2>\s*mg/dL、MAE\s*為\s*<p2000_mae:N2>\s*mg/dL，"
        r"持續性基準\s*RMSE\s*為\s*<p2000_base:N2>\s*mg/dL，相對改善約\s*<p2000_skill:I>%"),
    Pat(r"改善幅度最小者為\s*<pool_skill_min:N1>%"),
    # 〔稽核修正 A-10〕¶833 與 Alredaini 等(2026)之比較
    Pat(r"較該文\s*LSTM\s*之\s*13\.38\s*與\s*23\.94\s*mg/dL\s*分別低\s*<alr_lstm30_rel:N1>%\s*與\s*<alr_lstm60_rel:N1>%"),
    Pat(r"本研究分別高出\s*<alr_gpt30_diff:N2>\s*與\s*<alr_gpt60_diff:N2>\s*mg/dL"),
    Pat(r"於\s*30\s*與\s*60\s*分鐘<claim_gpt_both:G>本研究"),
    # 〔稽核修正 A-05〕App 內 30 分鐘時界之文字(const P.lstm.phNote 與「參數稽核」表),重跑後由 更新App常數與內嵌圖.py 一併更新
    Pat(r"30\s*分鐘時界下\s*RMSE\s*平均\s*<ph30_rmse_mean:N2>（中位數\s*<ph30_rmse_med:N2>）、MAE\s*平均\s*<ph30_mae_mean:N2>、"
        r"改善中位數\s*<ph30_skill_med:N1>%、勝率\s*<ph30_beat_pct:N1>%", dv={"App", "App2"}),
    Pat(r"30\s*分時界之改善中位數\s*<ph30_skill_med:N1>%", dv={"App", "App2"}),
    Pat(r"15\s*分鐘時界，本研究之\s*<pool_rmse_mean:N2>"),
    Pat(r"<claim_pool_better_per:C>之點落於對角線上方"),
    # 〔稽核修正 A-27〕簡報第 22 張改寫後之措辭(正文與備註)
    Pat(r"<claim_pool_better_per:C>之點（<pool_better_per_n:I>\s*[/／]\s*<pool_n:I>）落於對角線上方"),
    Pat(r"<claim_pool_better_per:C>的點在對角線上方"),
    TBL(r"表 4-7-6", r"持續性基準", r"RMSE", r"^<base_rmse_mean:N2>\s*mg/dL$"),
    TBL(r"表 4-7-6", r"逐人專屬", r"RMSE", r"^<per_rmse_mean:N2>\s*mg/dL$"),
    TBL(r"表 4-7-6", r"逐人專屬", r"改善中位數", r"^<per_skill_med:N1>%$"),
    TBL(r"表 4-7-6", r"逐人專屬", r"勝過基準", r"^<per_beat_pct:N1>%$"),
    TBL(r"表 4-7-6", r"跨病患合併", r"訓練視窗數", r"^<pool_train_windows:NT>$"),
    TBL(r"表 4-7-6", r"跨病患合併", r"RMSE", r"^<pool_rmse_mean:N2>\s*mg/dL$"),
    TBL(r"表 4-7-6", r"跨病患合併", r"改善中位數", r"^<pool_skill_med:N1>%$"),
    TBL(r"表 4-7-6", r"跨病患合併", r"勝過基準", r"^<pool_beat_pct:N1>%$"),
    TBL(r"簡報表", r"(?:技巧分數|改善)中位數", r"合併訓練", r"^<pool_skill_med:N1>\s*%$"),   # 〔2026-09-29 用詞統一〕列名接受兩種寫法
    TBL(r"簡報表", r"(?:技巧分數|改善)中位數", r"逐人專屬", r"^<per_skill_med:N1>\s*%$"),
    TBL(r"簡報表", r"勝過持續性基準", r"合併訓練", r"^<pool_beat_pct:N1>\s*%$"),
    TBL(r"簡報表", r"RMSE\s*中位數", r"合併訓練", r"^<pool_rmse_med:N2>\s*mg/dL$"),
    TBL(r"簡報表", r"RMSE\s*中位數", r"逐人專屬", r"^基準\s*<base_rmse_med:N2>\s*mg/dL$"),
    # ---------------- LSTM:30、60 分鐘與彙總
    Pat(r"30\s*分鐘時界，合併訓練之\s*RMSE\s*為\s*<ph30_rmse_mean:N2>\s*mg/dL、MAE\s*<ph30_mae_mean:N2>\s*mg/dL、"
        r"改善中位數\s*<ph30_skill_med:N1>%、勝率\s*<ph30_beat_pct:N1>%"),
    Pat(r"60\s*分鐘時界則分別為\s*<ph60_rmse_mean:N2>、<ph60_mae_mean:N2>、<ph60_skill_med:N1>%\s*與\s*<ph60_beat_pct:N1>%"),
    Pat(r"30\s*分鐘時分別為\s*<ph30_base_rmse_mean:N2>、<ph30_rmse_mean:N2>（MAE\s*<ph30_mae_mean:N2>）、"
        r"<ph30_skill_med:N1>%\s*與\s*<ph30_beat_pct:N1>%"),
    Pat(r"60\s*分鐘時為\s*<ph60_base_rmse_mean:N2>、<ph60_rmse_mean:N2>、<ph60_skill_med:N1>%\s*與\s*<ph60_beat_pct:N1>%"),
    Pat(r"30\s*分鐘時界之\s*MAE\s*<ph30_mae_mean:N2>\s*mg/dL"),
    Pat(r"30\s*分鐘均方根誤差\s*<ph30_base_rmse_mean:N2>、60\s*分鐘\s*<ph60_base_rmse_mean:N2>\s*mg/dL"),
    Pat(r"60\s*分鐘時界，本研究之\s*<ph60_rmse_mean:N2>"),
    Pat(r"本研究之\s*<ph60_rmse_mean:N2>\s*恰落於"),
    Pat(r"30\s*分鐘為\s*<ph30_pooled_rmse:N2>\s*mg/dL、60\s*分鐘為\s*<ph60_pooled_rmse:N2>\s*mg/dL"),
    Pat(r"本研究所得為\s*30\s*分鐘\s*<ph30_r2:N3>、60\s*分鐘\s*<ph60_r2:N3>"),
    Pat(r"30\s*分鐘時界之測試視窗為\s*<ph30_nwin:NT>\s*個"),
    Pat(r"30\s*分鐘時界\s*<ph30_nwin:NT>\s*個測試視窗"),
    Pat(r"三十分鐘時界下，(?:改善|技巧分數)中位數是\s*<ph30_skill_med:N1>%、MAE\s*<ph30_mae_mean:N2>"),   # 〔2026-09-29 用詞統一〕
    Pat(r"30\s*分時界為\s*<ph30_skill_med:N1>%"),
    Pat(r"30\s*分鐘時界下\s*RMSE\s*<ph30_rmse_mean:N2>、MAE\s*<ph30_mae_mean:N2>、改善中位數\s*<ph30_skill_med:N1>%、勝率\s*<ph30_beat_pct:N1>%"),
    TBL(r"表 5-4-2", r"^15\s*分鐘", r"本研究", r"^<pool_rmse_mean:N2>$"),
    TBL(r"表 5-4-2", r"^60\s*分鐘", r"本研究", r"^<ph60_rmse_mean:N2>$"),
    TBL(r"表 5-4-3", r"LSTM RMSE（30\s*分）", r"本研究", r"^<ph30_pooled_rmse:N2>\s*mg/dL（逐紀錄平均\s*<ph30_rmse_mean:N2>）"),
    TBL(r"表 5-4-3", r"LSTM RMSE（60\s*分）", r"本研究", r"^<ph60_pooled_rmse:N2>\s*mg/dL（逐紀錄平均\s*<ph60_rmse_mean:N2>）"),
    TBL(r"表 5-4-3", r"決定係數", r"本研究", r"^<ph30_r2:N3>／<ph60_r2:N3>"),
    # ---------------- 族群落差、併發症分組、長度分組、相關
    Pat(r"落差由\s*<gap_per:N1>\s*個百分點縮(?:小)?至\s*<gap_pool:N1>\s*個百分點"),
    Pat(r"from <gap_per:N1> to <gap_pool:N1> percentage points"),
    Pat(r"兩組落差達\s*<gap_per:N1>\s*個百分點"),
    Pat(r"落差由\s*<gap_per:N1>\s*收斂至\s*<gap_pool:N1>\s*個百分點"),
    Pat(r"落差從\s*<gap_per:N1>\s*個百分點收斂到\s*<gap_pool:N1>"),
    Pat(r"表現最差（改善中位數\s*<per_skill_comp1:N1>%(?:、勝率\s*<per_beat_comp1:N1>%)?）"),
    Pat(r"表現最佳（<per_skill_comp0:N1>%、勝率\s*<per_beat_comp0:N1>%）"),
    Pat(r"有併發症者平均\s*<npoints_comp1:NT>\s*點、無併發症者\s*<npoints_comp0:NT>\s*點"),
    Pat(r"合併為單一訓練集，共\s*<pool_train_windows:NT>\s*個視窗"),
    Pat(r"<pool_n:I>\s*筆中僅餘\s*<pool_nlose:I>\s*筆不如基準"),
    Pat(r"AUC\s*由\s*<sens24_micro_med:N3>\s*升至\s*<sens10_micro_med:N3>，(?:95%\s*|重複間)區間寬度由\s*<sens24_micro_w:N3>\s*縮為\s*<sens10_micro_w:N3>"),
    Pat(r"有併發症者之平均監測點數為\s*<npoints_comp1:NT>，無併發症者為\s*<npoints_comp0:NT>"),
    Pat(r"監測長度與糖化血色素之相關為\s*<corr_hba1c:N3>"),
    Pat(r"關聯由\s*<corr_zero:N2>\s*降至\s*<partial_r:N2>"),
    Pat(r"改以其對數控制為\s*<partial_log:N2>、等級偏相關為\s*<partial_rank:N2>"),
    Pat(r"長度高度相關（r\s*≈\s*<r_len_per:N2>）"),
    Pat(r"相關由\s*r\s*≈\s*<r_len_per:N2>\s*降至\s*<r_len_pool:N2>"),
    Pat(r"最短四分之一者之改善中位數為\s*<lg1_per_skill:N1>%、勝率僅\s*<lg1_per_beat:N1>%，最長四分之一者則為\s*"
        r"<lg4_per_skill:N1>%、勝率\s*<lg4_per_beat:N1>%"),
    Pat(r"最短四分之一組之改善中位數由\s*<lg1_per_skill:N1>%\s*轉為\s*<lg1_pool_skill:N1>%、勝率由\s*<lg1_per_beat:N1>%\s*升至\s*<lg1_pool_beat:N1>%"),
    Pat(r"表現趨於一致（<lg_pool_skill_min:N1>%\s*[–—-]\s*<lg_pool_skill_max:N1>%）"),
] + [
    TBL(r"表 4-7-7", rl, cl, rf"^<{k}{i}:N1>%$")
    for rl, i in [(r"無併發症", "comp0"), (r"有併發症", "comp1")]
    for cl, k in [(r"逐人模型改善", "per_skill_"), (r"合併訓練改善", "pool_skill_"),
                  (r"逐人勝率", "per_beat_"), (r"合併勝率", "pool_beat_")]
] + [
    TBL(r"表 4-7-7", rl, cl, rf"^<lg{i}_{k}:N1>%$")
    for rl, i in [(r"^最短\s*25%", 1), (r"^短（", 2), (r"^長（", 3), (r"^最長\s*25%", 4)]
    for cl, k in [(r"逐人模型改善", "per_skill"), (r"合併訓練改善", "pool_skill"),
                  (r"逐人勝率", "per_beat"), (r"合併勝率", "pool_beat")]
] + [
    TBL(r"表 4-7-7", r"", r"^分組$", rf"^{lab}（平均\s*<lg{i}_mean_points:NT>\s*點）$")
    for lab, i in [(r"最短\s*25%", 1), (r"短", 2), (r"長", 3), (r"最長\s*25%", 4)]
] + [
    # ---------------- 四象限(紀錄層級,quadrant_summary.csv)
    Pat(r"四象限之人數分布[^（]{0,8}（<qn_A:I>、<qn_B:I>、<qn_C:I>、<qn_D:I>）"),
    Pat(r"（A\s*<qn_A:I>、B\s*<qn_B:I>、C\s*<qn_C:I>、D\s*<qn_D:I>）"),
    Pat(r"實際併發症比例由\s*A\s*至\s*D\s*呈遞減（<qcomp_A:N1>%\s*→\s*<qcomp_B:N1>%\s*→\s*<qcomp_C:N1>%\s*→\s*<qcomp_D:N1>%）"),
    Pat(r"A\s*象限\s*<qcomp_A:N1>%\s*遞減至\s*D\s*象限\s*<qcomp_D:N1>%"),
    Pat(r"A象限實際併發症<qcomp_A:N1>%，遞減到D象限<qcomp_D:N1>%"),
    Pat(r"實際併發症比例(?:自|從)\s*<qcomp_A:N1>%\s*遞減(?:至|到)\s*<qcomp_D:N1>%"),
    Pat(r"實際併發症比例\s*<qcomp_A:N1>%\s*→\s*<qcomp_D:N1>%"),
    Pat(r"C\s*象限雖長期風險低（<qlong_C:N1>%），其短期風險卻高於\s*A\s*象限（<qshort_C:N1>%\s*對\s*<qshort_A:N1>%）"),
    Pat(r"C\s*象限之存在（長期風險\s*<qlong_C:N1>%、短期風險\s*<qshort_C:N1>%）"),
    Pat(r"長期風險幾乎相同（<qlong_A:N1>%\s*與\s*<qlong_B:N1>%"),
    Pat(r"幾乎相同[—–-]+<qlong_A:N1>對<qlong_B:N1>"),
] + [
    x for k, (L, rrow, card) in _Q.items() for x in [
        TBL(r"表 4-8-1", rrow, r"人數", rf"^<qn_{k}:I>$"),
        TBL(r"表 4-8-1", rrow, r"平均長期風險", rf"^<qlong_{k}:N1>%$"),
        TBL(r"表 4-8-1", rrow, r"平均短期風險", rf"^<qshort_{k}:N1>%$"),
        TBL(r"表 4-8-1", rrow, r"實際併發症比例", rf"^<qcomp_{k}:N1>%$"),
        TBL(r"簡報表", rrow, r"人數", rf"^<qn_{k}:I>$"),
        TBL(r"簡報表", rrow, r"平均長期風險", rf"^<qlong_{k}:N1>%$"),
        TBL(r"簡報表", rrow, r"平均短期風險", rf"^<qshort_{k}:N1>%$"),
        TBL(r"簡報表", rrow, r"實際併發症比例", rf"^<qcomp_{k}:N1>%$"),
        Pat(card + rf"｜<qn_{k}:I>\s*筆紀錄"),
        Pat(rf"併發症\s*<qcomp_{k}:N1>%・成本\s*<qcost_{k}:N1>\s*萬", kind={"框"}, prv=card),
        Pat(rf"實際併發症比例\s*<qcomp_{k}:N1>%、餘生費用平均約\s*<qcost_{k}:N1>\s*萬（95%\s*模擬區間\s*"
            rf"<qlo_{k}:N1>\s*[–—-]\s*<qhi_{k}:N1>\s*萬）", dv={"App"}, prv=rf"\b{L}:\{{name"),
        # 成本(病患層級,quadrant_cost_interval.csv)
        TBL(r"表 4-9-3", rrow, r"人數", rf"^<qcn_{k}:I>$"),
        TBL(r"表 4-9-3", rrow, r"平均", rf"^<qcost_{k}:N1>$"),
        TBL(r"表 4-9-3", rrow, r"區間", rf"^<qlo_{k}:N1>\s*[–—-]\s*<qhi_{k}:N1>$"),
    ]
] + [
    # ---------------- 成本
    Pat(r"長期風險低者(?:（C、D）)?之\s*<cost_ratio_ABCD:N2>\s*倍"),
    Pat(r"本研究所報告之\s*<cost_ratio_ABCD:N2>\s*倍"),
    Pat(r"approximately <cost_ratio_ABCD:N2> times the lifetime cost"),
    Pat(r"約為低者之\s*<cost_ratio_ABCD:N2>\s*倍"),
    Pat(r"^\s*<cost_ratio_ABCD:N2>\s*倍\s*$", kind={"框"}, nxt=r"長期風險"),
    Pat(r"A\s*與\s*B\s*象限成本相近（<qcost_A:N1>\s*與\s*<qcost_B:N1>\s*萬元）"),
    Pat(r"倍（<qcost_A:N1>\s*與\s*<qcost_B:N1>\s*萬對\s*<qcost_C:N1>\s*與\s*<qcost_D:N1>\s*萬元）"),
    Pat(r"A\s*<qcost_A:N1>、B\s*<qcost_B:N1>、C\s*<qcost_C:N1>、D\s*<qcost_D:N1>\s*萬"),
    Pat(r"A約<qcost_A:N1>萬、B\s*<qcost_B:N1>、C\s*<qcost_C:N1>、D\s*<qcost_D:N1>"),
    Pat(r"A、B\s*兩象限（<qlo_A:N1>\s*[–—-]\s*<qhi_A:N1>\s*與\s*<qlo_B:N1>\s*[–—-]\s*<qhi_B:N1>\s*萬元）"),
    Pat(r"C、D\s*兩象限（<qlo_C:N1>\s*[–—-]\s*<qhi_C:N1>\s*與\s*<qlo_D:N1>\s*[–—-]\s*<qhi_D:N1>\s*萬元）"),
    Pat(r"每人餘生醫療成本平均為\s*<cost_mean:NT>\s*元、中位數\s*<cost_median:NT>\s*元"),
    Pat(r"若以\s*109\s*筆紀錄為分母則為\s*<cost_mean_rec:NT>\s*元"),
    Pat(r"母體餘生總成本（100\s*位病患）平均約\s*<pop_total_wan:NT>\s*萬元，95%\s*模擬區間為\s*<pop_lo_wan:NT>\s*萬至\s*<pop_hi_wan:NT>\s*萬元"),
    TBL(r"表 4-9-4", r"實際盛行率", r"大血管", r"^<macro_prev:N1>%$"),
    TBL(r"表 4-9-4", r"實際盛行率", r"小血管", r"^<micro_prev:N1>%$"),
    TBL(r"表 4-9-4", r"balanced", r"大血管預測", r"^<macro_pred_bal:N1>%$"),
    TBL(r"表 4-9-4", r"balanced", r"小血管預測", r"^<micro_pred_bal:N1>%$"),
    TBL(r"表 4-9-4", r"未加權", r"大血管預測", r"^<macro_pred_nw:N1>%$"),
    TBL(r"表 4-9-4", r"未加權", r"小血管預測", r"^<micro_pred_nw:N1>%$"),
    TBL(r"表 4-9-4", r"balanced", r"AUC", r"^<macro_auc_bal:N3>\s*/\s*<micro_auc_bal:N3>$"),
    TBL(r"表 4-9-4", r"未加權", r"AUC", r"^<macro_auc_nw:N3>\s*/\s*<micro_auc_nw:N3>$"),
    TBL(r"表 4-9-4", r"balanced", r"餘生(?:醫療)?成本", r"每人\s*<cost_mean_wan:N1>\s*萬、母體\s*<pop_total_wan:NT>\s*萬"),
    TBL(r"表 4-9-4", r"未加權", r"餘生(?:醫療)?成本", r"每人\s*<cost_mean_wan_nw:N1>\s*萬、母體\s*<pop_total_wan_nw:NT>\s*萬"),
    Pat(r"成本中位數為\s*<cost_median_wan:N1>\s*萬元，第\s*75\s*百分位\s*<cost_p75_wan:N1>\s*萬元，第\s*90\s*百分位\s*"
        r"<cost_p90_wan:N1>\s*萬元，最高達\s*<cost_max_wan:N1>\s*萬元"),
    Pat(r"中位數\s*<cost_median_wan:N1>\s*萬、第\s*90\s*百分位\s*<cost_p90_wan:N1>\s*萬、最高\s*<cost_max_wan:N1>\s*萬"),
    Pat(r"中位數<cost_median_wan:N1>萬，但第90百分位<cost_p90_wan:N1>萬、最高<cost_max_wan:N1>萬"),
    Pat(r"前\s*20%\s*(?:病患)?\s*(?:占|佔)(?:全體(?:餘生總成本之|支出之?)?)?\s*<top20:N1>%"),
    Pat(r"前\s*10%\s*病患占全體(?:餘生總成本之|支出)\s*<top10:N1>%"),
    Pat(r"前\s*30%\s*占\s*<top30:N1>%，前\s*50%\s*占\s*<top50:N1>%"),
    Pat(r"accounting for <top20:N1>% of total expenditure"),
    Pat(r"Gini\s*(?:係數|coefficient|=)\s*(?:為\s*)?<gini:N3>"),
    # ---------------- 馬可夫與 CGM
    Pat(r"\|λ₂\|\s*=\s*<lambda2_3:N3>"),
    Pat(r"收斂半衰期僅約\s*\d+\s*步（<halflife3_h:N1>\s*小時）"),
    Pat(r"半衰期\s*<halflife5_h:N1>\s*小時對\s*<halflife3_h:N1>\s*小時"),
    Pat(r"目標範圍內比例（<st5_in:N1>%）與三態結果（<st3_in:N1>%）"),
    Pat(r"低\s*<st3_low:N1>%、範圍內\s*<st3_in:N1>%、高\s*<st3_high:N1>%"),
    TBL(r"表 4-7-3", r"三態（穩態", r"低血糖", r"^<st3_low:N1>%$"),
    TBL(r"表 4-7-3", r"三態（穩態", r"目標範圍內", r"^<st3_in:N1>%$"),
    TBL(r"表 4-7-3", r"三態（穩態", r"高血糖", r"^<st3_high:N1>%$"),
    TBL(r"表 4-7-3", r"三態（穩態", r"λ", r"^<lambda2_3:N3>$"),
    TBL(r"表 4-7-3", r"三態（穩態", r"半衰期", r"^<halflife3_h:N1>\s*小時$"),
    Pat(r"TIR\s*中位數為\s*<tir_med:N1>%"),
    Pat(r"變異係數的中位數為\s*<cv_med:N1>%"),
    Pat(r"範圍內<tir_med:N1>%、變異係數<cv_med:N1>%"),
    TBL(r"簡報表", r"TIR", r"數值", r"^<tir_med:N1>\s*%$"),
    TBL(r"簡報表", r"TAR", r"數值", r"^<tar_med:N1>\s*%$"),
    TBL(r"簡報表", r"TBR", r"數值", r"^<tbr_med:N1>\s*%$"),
    TBL(r"簡報表", r"變異係數", r"數值", r"^<cv_med:N1>\s*%$"),
    TBL(r"簡報表", r"GMI", r"數值", r"^<gmi_med:N2>\s*%$"),
] + [
    TBL(cap, rf"^{rname}(?:\s/\s|$)", rf"^{cname}$", rf"^<m3_{i}{j}:N3>$")
    for cap, rows, cols in [(r"表 4-7-2", ["低血糖", "目標範圍內", "高血糖"], ["低血糖", "目標範圍內", "高血糖"]),
                            (r"簡報表", ["低血糖", "範圍內", "高血糖"], ["低血糖", "範圍內", "高血糖"])]
    for i, rname in enumerate(rows) for j, cname in enumerate(cols)
] + [
    TBL(r"表 4-7-5", rf"^{rl}", rf"^{cl}", rf"^<m5_{i}{j}:N3>$")
    for i, rl in enumerate(["極低", "低（", "目標", "高（", "極高"])
    for j, cl in enumerate(["極低", "低 ", "目標", "高 ", "極高"])
] + [
    Pat(r"目標範圍內→目標範圍內\s*=\s*<m3_11:N3>"),
    Pat(r"除「極低血糖→目標範圍」（<m5_02:N3>）"),
    Pat(r"極低血糖之自我維持機率高達\s*<m5_00:N3>，明顯高於輕度低血糖之\s*<m5_11:N3>"),
    # 分群(cluster_complication_rate.csv:KMeans k=2)
    Pat(r"群\s*0\s*大血管併發症比例\s*<clu_macro0:N1>%、群\s*1\s*為\s*<clu_macro1:N1>%，任一併發症比例則為\s*<clu_any0:N1>%\s*對\s*<clu_any1:N1>%"),
    Pat(r"群間比例\s*<clu_macro0:N1>%\s*對\s*<clu_macro1:N1>%"),
    # 系統門檻(diagnostics:insample_median、flagged_pct_old_cut)
    Pat(r"預測中位數」（<insample_med_pct:N1>%）"),
    Pat(r"全資料配適之預測中位數（<insample_med_pct:N1>%）"),
    Pat(r"判為(?:長期)?高風險者達\s*<flagged_old_pct:N1>%（設計上應為\s*50%）"),
    Pat(r"大血管併發症之實際盛行率為\s*<macro_prev:N1>%，該模型之平均預測機率為\s*<macro_pred_bal:N1>%（高\s*<macro_pred_gap:N1>\s*個百分點）"),
    Pat(r"小血管併發症實際為\s*<micro_prev:N1>%，預測均值為\s*<micro_pred_bal:N1>%（高\s*<micro_pred_gap:N1>\s*個百分點）"),
    Pat(r"高風險比例為\s*<flagged_old_pct:N1>%（設計上應為\s*50%）"),
] + [
    # ---------------- 象限切點
    Pat(r"out-of-fold\s*預測中位數（<cut_long_pct:N1>%）"),
    Pat(r"長期中位數\s*<cut_long_pct:N1>%"),
    Pat(r"長期\s*<cut_long_pct:N1>%、短期\s*<cut_short_pct:N2>%"),
    Pat(r"交叉驗證預測之中位數\s*<cut_long_pct:N1>%、短期＝<cut_short_pct:N2>%"),
    # ---------------- 勝算比
    Pat(r"罹病年數（<or_micro_Duration:N2>，小血管）、空腹胰島素（<or_any_Ins_f:N2>，任一）、年齡（<or_macro_Age:N2>，大血管）"),
    Pat(r"罹病年數<or_micro_Duration:N2>、空腹胰島素<or_any_Ins_f:N2>、年齡<or_macro_Age:N2>"),
    Pat(r"高密度脂蛋白（<or_macro_HDL:N2>）與空腹\s*C-?胜肽（<or_micro_Cpep_f:N2>）"),
    Pat(r"高密度脂蛋白<or_macro_HDL:N2>與C胜肽<or_micro_Cpep_f:N2>"),
    Pat(r"尿酸勝算比\s*<or_any_UA:N2>"),
    Pat(r"小血管併發症最強之(?:單一)?預測因子（OR\s*≈\s*<or_micro_Duration:N2>）"),
    Pat(r"罹病時間對小血管併發症最強（OR\s*≈\s*<or_micro_Duration:N2>）、年齡對大血管併發症最強（OR\s*≈\s*<or_macro_Age:N2>）"),
    Pat(r"大血管併發症最強因子（OR\s*≈\s*<or_macro_Age:N2>）"),
    Pat(r"高密度脂蛋白對大血管具保護作用（OR\s*≈\s*<or_macro_HDL:N2>）"),
    Pat(r"較低之小血管風險相關（OR\s*≈\s*<or_micro_Cpep_f:N2>）"),
    # ---------------- 〔稽核修正 B-10〕共線性(¶521、¶574、¶585、¶802、¶803)
    Pat(r"估算腎絲球過濾率與肌酸酐（r\s*=\s*<r_eGFR_Cr:N2>）"),
    Pat(r"糖化血色素與糖化白蛋白（r\s*=\s*<r_HbA1c_GA:N2>）"),
    Pat(r"總膽固醇與低密度脂蛋白（r\s*=\s*<r_TC_LDL:N2>）"),
    Pat(r"空腹與飯後血糖（r\s*=\s*<r_FPG_PPG:N2>）"),
    Pat(r"空腹與飯後胰島素（r\s*=\s*<r_Ins_f_Ins_2h:N2>）"),
    Pat(r"近似指標糖化白蛋白（二者\s*r\s*[≈=]\s*<r_HbA1c_GA:N2>）"),
    Pat(r"體重與身體質量指數為\s*<r_Weight_BMI:N2>、身體質量指數與身高僅\s*<r_BMI_Height:N2>"),
    Pat(r"與肌酸酐\s*r\s*=\s*<r_UA_Cr:N2>、與估算腎絲球過濾率\s*r\s*=\s*<r_UA_eGFR:N2>"),
    Pat(r"相關係數\s*<bmi_def_r:N6>[，、,]\s*最大偏差\s*<bmi_def_dev:N3>"),
    Pat(r"體重、身體質量指數與身高分別為\s*<vif0_Weight:I>、<vif0_BMI:I>\s*與\s*<vif0_Height:I>"),
    # A-02:「前述各組之最大值（總膽固醇 9.3）」—— 期望值為五組變數之最大 VIF(實為肌酸酐),不是總膽固醇之 VIF
    Pat(r"前述各組之最大值（[^\d（）]{1,12}?\s*<vif1_maxpairs:N1>）"),
    # 依稽核建議改寫後之措辭(選用樣式;未改寫前找不到不算錯)
    Pat(r"肌酸酐與估算腎絲球過濾率之\s*(?:VIF|變異數膨脹因子)\s*亦逾\s*10\s*(?:倍\s*)?（<vif1_Cr:N1>\s*與\s*<vif1_eGFR:N1>）"),
    Pat(r"肌酸酐\s*<vif1_Cr:N1>、估算腎絲球過濾率\s*<vif1_eGFR:N1>"),
    Pat(r"判定係數依序為\s*<vifr2_Weight:N3>、<vifr2_BMI:N3>\s*與\s*<vifr2_Height:N3>"),
    Pat(r"變異數膨脹因子為體重\s*<vif0_Weight:I>、身體質量指數\s*<vif0_BMI:I>、身高\s*<vif0_Height:I>"),
    Pat(r"以\s*<boot_n:I>\s*次拔靴重抽"),
    Pat(r"少數方向之次數分別達\s*<boot_min_Weight:I>、<boot_min_BMI:I>\s*與\s*<boot_min_Height:I>\s*次"),
    Pat(r"未逾\s*5\s*倍之\s*<vif_n_lt5:I>\s*個特徵其中位數僅\s*<boot_med_lt5:N1>\s*次"),
    Pat(r"於\s*<abl_n:I>\s*組隨機切分下重算"),
    Pat(r"全部\s*24\s*個特徵之\s*AUC\s*中位數為\s*<abl_all:N3>；剔除身體質量指數後為\s*<abl_noBMI:N3>，"
        r"剔除身高與體重後為\s*<abl_noHW:N3>，再剔除估算腎絲球過濾率後為\s*<abl_noHWe:N3>"),
]

# 「九成」類主張:(判斷函式, 建議改寫)
CLAIMS = {
    "逾九成": lambda v: v > 90, "超過九成": lambda v: v > 90, "九成以上": lambda v: v >= 90,
    "近九成": lambda v: 85 <= v < 90, "約九成": lambda v: 85 <= v <= 95,
    # 〔稽核修正 A-27〕倍數主張(v 為倍數)
    "逾倍": lambda v: v > 2, "超過一倍": lambda v: v > 2, "逾一倍": lambda v: v > 2, "倍增": lambda v: v >= 1.95,
    # 〔稽核修正 A-10〕v 為「本研究高出 GPT-4.1」兩時界之較小者;須 > 0 才能說該文兩時界皆較佳
    "皆略優於": lambda v: v > 0, "皆優於": lambda v: v > 0, "均略優於": lambda v: v > 0, "均優於": lambda v: v > 0,
}
CLAIM_KEY = {"claim_pool_better_per": "pool_better_per_pct", "claim_skill_ratio": "pool_per_skill_ratio",
             "claim_gpt_both": "alr_gpt_min_diff"}
# 主張之期望值顯示方式:(位數, 單位, 措辭範例)
CLAIM_FMT = {"pool_better_per_pct": (1, "%", "約九成"), "pool_per_skill_ratio": (2, " 倍", "提升至約 N 倍/明顯提升"),
             "alr_gpt_min_diff": (2, " mg/dL", "依時界分述何者較佳")}

# App 常數區 const P 逐鍵比對:(路徑, 期望值之鍵)
APP_CHECKS = (
    [(f"cohort.auc.{t}", f"auc_{t}_lr") for t in ("any", "micro", "macro")]
    + [(f"cohort.aucCI.{t}[{i}]", f"ci_{t}_lr_{s}") for t in ("any", "micro", "macro") for i, s in ((0, "lo"), (1, "hi"))]
    + [(f"cohort.quadInfo.{k}.n", f"qn_{k}") for k in "ABCD"]
    + [(f"cohort.quadInfo.{k}.real", f"qcomp_{k}") for k in "ABCD"]
    + [(f"cohort.quadCost.{k}", f"qcost_raw_{k}") for k in "ABCD"]
    + [("cohort.costMedian", "cost_median"), ("cuts.longMed", "cut_long_raw"), ("cuts.shortMed", "cut_short_raw"),
       ("lstm.skillMedian", "pool_skill_med"), ("lstm.beatPct", "pool_beat_pct"), ("lstm.rmseMedian", "pool_rmse_med"),
       ("lstm.baseRmseMedian", "base_rmse_med"), ("lstm.perSkillMedian", "per_skill_med"), ("lstm.n", "pool_n"),
       ("lstm.rmseMean", "pool_rmse_mean"), ("lstm.baseRmseMean", "base_rmse_mean"), ("cost.infl", "cpi_cum")]
    + [(f"cost.{k}", f"cost_direct_{k}") for k in ("none", "macro", "micro", "both")]
    + [(f"markov.M[{i}][{j}]", f"m3_{i}{j}") for i in range(3) for j in range(3)]
    + [(f"markov.stationary[{i}]", f"st3_frac{i}") for i in range(3)]
    + [(f"cgmCohort.{c}.{s}", f"{k}_{s}_raw") for c, k in (("TIR", "tir"), ("TAR", "tar"), ("TBR", "tbr"), ("CV", "cv"), ("GMI", "gmi"))
       for s in ("mean", "p25", "p50", "p75")]
    + [(f"cgmCohort.mean.{s}", f"glu_{s}_raw") for s in ("mean", "p25", "p50", "p75")]
)


def _const_block(s, name):
    m = re.search(r"\bconst\s+" + name + r"\s*=\s*\{", s)
    if not m:
        return None, None
    j, d = m.end() - 1, 0
    for k in range(j, len(s)):
        if s[k] == "{":
            d += 1
        elif s[k] == "}":
            d -= 1
            if d == 0:
                return s[j:k + 1], s[:m.start()].count("\n") + 1
    return None, None


def _get(obj, path):
    for part in re.findall(r"[^.\[\]]+|\[\d+\]", path):
        if part.startswith("["):
            i = int(part[1:-1])
            if not isinstance(obj, list) or i >= len(obj):
                return None
            obj = obj[i]
        else:
            if not isinstance(obj, dict) or part not in obj:
                return None
            obj = obj[part]
    return obj


def _decimals(x):
    if isinstance(x, (int, np.integer)) or float(x).is_integer():
        return 0
    t = repr(float(x))
    if "e" in t:
        return 15
    return len(t.split(".")[1]) if "." in t else 0


# ======================================================================
# 〔稽核修正 CODE-29〕豁免清單:鍵＝(檔, 值, 前後文);前後文為正規式,須出現在該處前後 40 字
#   或該格之表題/列/欄中。不再以(檔,值)對整檔一體豁免。
# ======================================================================
def load_allow(path, tool):
    rows = []
    if not os.path.exists(path):
        return rows
    import io
    raw = open(path, "rb").read()
    try:
        txt = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        # 〔稽核修正 CODE-29〕Windows 上以 Excel「CSV(逗號分隔)」另存會變成 cp950;改以 cp950 讀取並提示
        txt = raw.decode("cp950", errors="replace")
        print(f"[allow] {os.path.basename(path)} 不是 UTF-8(可能由 Excel 另存為「CSV(逗號分隔)」),"
              f"已改以 cp950 讀取;建議另存為「CSV UTF-8(逗號分隔)」。")
    # 「#」開頭為註解(Excel 另存後可能變成 "# …" 帶引號,一併視為註解)
    txt = "\n".join(l for l in txt.splitlines() if not l.lstrip().lstrip('"').startswith("#"))
    df = pd.read_csv(io.StringIO(txt), dtype=str).fillna("")
    for _, r in df.iterrows():
        use = (r.get("適用", "") or "*").strip()
        if use not in ("*", tool):
            continue
        ctx = (r.get("前後文", "") or "").strip()
        if not ctx:
            print(f"[allow] 略過未寫前後文之豁免列:{r.get('檔')},{r.get('值')}(豁免必須指明前後文)")
            continue
        try:
            cx = re.compile(ctx)
        except re.error as e:                          # 〔稽核修正 CODE-29〕前後文寫錯時略過該列並指出原因
            print(f"[allow] 略過前後文不是有效正規式之豁免列:{r.get('檔')},{r.get('值')},{ctx!r}({e})")
            continue
        rows.append((r["檔"].strip(), r["值"].strip(), cx, r.get("理由", "").strip()))
    return rows


def allowed(rows, dv, tok, window, ctx):
    for f, v, cx, why in rows:
        if v != tok:
            continue
        if f not in ("*", dv) and not (f == "App-常數區" and dv == "App"):
            continue
        if cx.search(window) or cx.search(ctx):
            return why
    return None


def snippet(raw, a, b, w=20):
    return (raw[max(0, a - w):a] + "【" + raw[a:b] + "】" + raw[b:b + w]).replace("\n", "⏎")


# ======================================================================
def main():
    ap = argparse.ArgumentParser(description="第 2 層:交付物數字逐處比對")
    ap.add_argument("--base", default="", help="交付物目錄(含 專題書面報告_v*.docx)")
    ap.add_argument("--out", default="", help="「程式與輸出」目錄(含 output_risk/ 等)")
    ap.add_argument("--allow", default="", help="豁免清單,預設為同目錄之 allowlist.csv")
    ap.add_argument("--mismatch-csv", default="", help="不符清單(工作清單)輸出路徑")
    ap.add_argument("--show-ok", action="store_true", help="另列出每一處相符之位置")
    ap.add_argument("--show-uncovered", action="store_true", help="另列出同值但未被樣式涵蓋之出現處")
    ap.add_argument("--show-source", action="store_true", help="(相容舊版)印出每個期望值之來源")
    a = ap.parse_args()

    BASE = _find(_is_deliv, a.base)
    OUT = _find(_is_out, a.out)
    if not BASE:
        sys.exit("[錯誤] 找不到交付物目錄(應含 專題書面報告_v*.docx)。請用 --base 指定。")
    if not OUT:
        sys.exit("[錯誤] 找不到「程式與輸出」目錄(應含 output_risk/quadrant_summary.csv)。請用 --out 指定。")
    print("【第 2 層】交付物數字逐處比對(每一處出現位置都必須等於本次輸出之值)")
    print(f"[path] 交付物:{BASE}")
    print(f"[path] 程式輸出:{OUT}")
    FILES = deliverables(BASE)
    for k, v in FILES.items():
        print(f"[file] {k}: {os.path.basename(v) if v else '(找不到)'}")
    miss = [k for k, v in FILES.items() if not v or not os.path.exists(v)]
    if miss:
        sys.exit(f"[錯誤] 找不到交付物:{miss}")

    V = derive(OUT)
    print(f"[期望值] 由本次輸出推導 {len(V)} 個量(ROUND_HALF_UP)")
    for sec, err in DERIVE_ERRORS:
        print(f"[期望值] ★ 區段「{sec}」推導失敗:{err}(該區段之項目將列為無法檢查)")
    ref_dir = os.path.join(OUT, "reference_teacher")
    VT = {}
    if os.path.isdir(ref_dir):
        try:
            VT = derive(ref_dir)
            print(f"[另一次執行] reference_teacher/:{len(VT)} 個量(僅用於標註「疑似混用另一次執行」;"
                  f"{len(DERIVE_ERRORS)} 個區段因舊版欄位不同而略過)")
        except Exception as e:
            print(f"[另一次執行] reference_teacher/ 無法推導:{e}")

    U = units_docx(FILES["報告"]) + units_pptx(FILES["簡報"], "簡報") + units_pptx(FILES["海報"], "海報")
    U += units_html(FILES["App"])
    same_app = open(FILES["App"], "rb").read() == open(FILES["App2"], "rb").read()
    print(f"[file] App 兩檔 {'逐位元相同(只掃描 index.html)' if same_app else '★不同★(兩檔皆掃描)'}")
    if not same_app:
        U += units_html(FILES["App2"], "App2")
    from collections import Counter
    print("[scope] 文字單位:" + "、".join(f"{k} {v}" for k, v in Counter(u.dv for u in U).items()))

    # ---- 逐處擷取
    occ, taken = [], {}
    for u in U:
        for p in PATTERNS:
            if not p.applies(u):
                continue
            for m in p.rx.finditer(u.text):
                for g, val in m.groupdict().items():
                    if val is None:
                        continue
                    key = g.split("__")[0]
                    span = m.span(g)
                    tk = (u.dv, u.loc, span)
                    if tk in taken:
                        continue
                    taken[tk] = key
                    occ.append((u, key, val, span))

    def normnum(s):
        return s.replace("+", "").strip()

    rows, stat = [], {}
    n_ok = n_bad = n_bd = 0
    ok_list = []
    for u, key, val, (s0, s1) in occ:
        st = stat.setdefault(key, {})
        cell = st.setdefault(u.dv, [0, 0, 0])
        ctxs = snippet(u.raw, s0, s1)
        if key in CLAIM_KEY:
            vk = CLAIM_KEY[key]
            if vk not in V:
                continue
            v = V[vk][0]
            good = CLAIMS.get(val, lambda x: True)(v)
            nd_, unit_, eg_ = CLAIM_FMT.get(vk, (1, "%", "約九成"))
            exp = f"{half_up(v, nd_)[0]}{unit_}"
            if good:
                cell[0] += 1
                n_ok += 1
            else:
                cell[1] += 1
                n_bad += 1
                rows.append(dict(交付物=u.dv, 位置=u.loc, 項目=NAMES.get(vk, vk), 前後文=ctxs, 現值=val,
                                 應為值=f"與 {exp} 相符之措辭(例:{eg_}/{exp})", 判定="不符(主張與數值不符)",
                                 來源=V[vk][3], 備註=""))
            continue
        if key not in V or V[key][1] is None:
            cell[2] += 0
            continue
        v, nd, thou, src = V[key]
        exp, bd = half_up(v, nd, thou)
        got = normnum(val)
        if got == exp and not bd:
            cell[0] += 1
            n_ok += 1
            ok_list.append((u, key, ctxs, got))
        elif bd and got in bd:
            # 〔稽核修正 A-19〕來源 CSV 已捨入且恰在進位邊界:即使與「再進位」結果相同也可能是
            #   雙重進位之錯(例:0.8875→0.888,全精度 0.887474→0.887),一律列為邊界待確認。
            cell[2] += 1
            n_bd += 1
            rows.append(dict(交付物=u.dv, 位置=u.loc, 項目=NAMES.get(key, key), 前後文=ctxs, 現值=got,
                             應為值=f"{bd[0]} 或 {bd[1]}", 判定="▲邊界(來源 CSV 已捨入,須以全精度確認)",
                             來源=f"{src}＝{v!r}", 備註=""))
        else:
            cell[1] += 1
            n_bad += 1
            note = ""
            if key in VT and VT[key][1] is not None and half_up(VT[key][0], nd, thou)[0] == got:
                note = "現值＝reference_teacher(另一次執行)之值:疑似混用兩次執行"
            rows.append(dict(交付物=u.dv, 位置=u.loc, 項目=NAMES.get(key, key), 前後文=ctxs, 現值=got,
                             應為值=exp, 判定="不符", 來源=src, 備註=note))

    # ---- App 常數區 const P
    app_rows = []
    constP_lines = set()                                 # const P 所在行(已逐鍵比對,殘留掃描略過)
    for fn in (["App"] if same_app else ["App", "App2"]):
        s = open(FILES[fn], encoding="utf-8", errors="replace").read()
        blk, line = _const_block(s, "P")
        if blk is not None:
            constP_lines |= {(fn, f"L{line + k}") for k in range(blk.count("\n") + 1)}
        if blk is None:
            app_rows.append((fn, "const P", "★ 找不到 const P"))
            n_bad += 1
            continue
        try:
            P = json.loads(blk)
        except Exception as e:
            app_rows.append((fn, "const P", f"★ JSON 解析失敗:{e}"))
            n_bad += 1
            continue
        n_app = 0
        for path, key in APP_CHECKS:
            got = _get(P, path)
            if got is None or key not in V or isinstance(got, (dict, list, str)):
                continue
            n_app += 1
            v = V[key][0]
            tol = 0.5 * 10 ** (-min(_decimals(got), _decimals(v))) + 1e-12
            st = stat.setdefault(key, {}).setdefault(fn, [0, 0, 0])
            if abs(float(got) - v) <= tol:
                st[0] += 1
                n_ok += 1
            else:
                st[1] += 1
                n_bad += 1
                rows.append(dict(交付物=fn, 位置=f"L{line} const P.{path}", 項目=NAMES.get(key, key),
                                 前後文=f"P.{path} = {got}", 現值=str(got), 應為值=f"{v:.10g}(容差 {tol:.1g})",
                                 判定="不符", 來源=V[key][3], 備註="App 常數"))
        app_rows.append((fn, f"L{line} const P", f"逐鍵比對 {n_app} 項"))

    # ---- 無法檢查(來源 CSV 缺)與找不到任何出現位置(樣式失效或已刪除)之項目
    pat_keys = {g.split("__")[0] for p in PATTERNS for g in p.rx.groupindex}
    occ_keys = {k for _u, k, _v, _s in occ}
    # 〔稽核修正 B-10〕選用樣式(OPTIONAL_KEYS)只有在交付物中確實出現時才要求可推導
    underivable = sorted(k for k in pat_keys if k not in V and k not in CLAIM_KEY
                         and (k not in OPTIONAL_KEYS or k in occ_keys))

    def need_file(k):
        for rx, fn in NEED_FILE:
            if re.search(rx, k):
                return fn
        return "推導此量所需之輸出 CSV"
    for k in underivable:
        rows.append(dict(交付物="(全部)", 位置="—", 項目=NAMES.get(k, k), 前後文="—", 現值="—", 應為值="—",
                         判定="無法檢查", 來源=need_file(k),
                         備註=f"{need_file(k)} 不存在或缺欄(未執行 run_all.py --with-sensitivity?)"))
    expected_keys = sorted((pat_keys & set(V)) | set(CLAIM_KEY))
    never = [k for k in expected_keys
             if sum(sum(c[:3]) for c in stat.get(k, {}).values()) == 0 and k not in CLAIM_KEY
             and k not in OPTIONAL_KEYS]
    optional_absent = [k for k in expected_keys if k in OPTIONAL_KEYS and k not in occ_keys]
    for k in never:
        rows.append(dict(交付物="(全部)", 位置="—", 項目=NAMES.get(k, k), 前後文="—", 現值="—",
                         應為值=half_up(V[k][0], V[k][1], V[k][2])[0], 判定="找不到任何出現位置",
                         來源=V[k][3], 備註="樣式可能因改寫而失效,或該數字已自交付物刪除;請確認後修改 PATTERNS"))

    # ---- 疑似殘留(reference_teacher 之值、未被任何樣式涵蓋)
    allow_path = a.allow or os.path.join(HERE, "allowlist.csv")
    ALLOW = load_allow(allow_path, "crosscheck")
    print(f"[allow] 豁免清單 {os.path.basename(allow_path)}:{len(ALLOW)} 列(須符合前後文才豁免)")
    covered = {}
    for u, key, val, span in occ:
        covered.setdefault((u.dv, u.loc), []).append(span)
    residual, exempt = [], []
    # 〔稽核修正 B-03〕只掃描「有鑑別力」之舊值:非整數、有效位數 ≥3、非整十/整百之值,
    #   且不與本次執行之任一值相同(相同者無法分辨是舊值還是另一個量)。
    current_strs = {half_up(v, nd, th)[0] for (v, nd, th, _s) in V.values() if nd is not None}
    for key, (vt, nd, thou, _src) in VT.items():
        if nd is None or nd == 0 or key not in V:
            continue
        old = half_up(vt, nd, thou)[0]
        new = half_up(V[key][0], nd, thou)[0]
        digits = re.sub(r"[^\d]", "", old).lstrip("0")
        if (old == new or len(digits) < 3 or old in current_strs
                or float(old.replace(",", "")) % 10 == 0):      # 0.0、50.0、100.0 等整十值無鑑別力
            continue
        rx = re.compile(r"(?<![\d.,])" + re.escape(old) + r"(?![\d])")
        for u in U:
            if (u.dv, u.loc) in constP_lines:
                continue
            for m in rx.finditer(u.text):
                if any(s0 <= m.start() < s1 for s0, s1 in covered.get((u.dv, u.loc), [])):
                    continue
                win = u.raw[max(0, m.start() - 40):m.end() + 40]
                why = allowed(ALLOW, u.dv, old, win, u.ctx)
                if why:
                    exempt.append((u.dv, u.loc, NAMES.get(key, key), old, why))
                else:
                    residual.append((u, key, old, new, m.span()))
    for u, key, old, new, (s0, s1) in residual:
        n_bad += 1
        rows.append(dict(交付物=u.dv, 位置=u.loc, 項目=NAMES.get(key, key), 前後文=snippet(u.raw, s0, s1),
                         現值=old, 應為值=f"{new}(若此處確指該量)", 判定="疑似殘留(另一次執行之值)",
                         來源="reference_teacher/", 備註="未被任何樣式涵蓋;若非該量,請加入 allowlist.csv 並寫明前後文"))

    # ---- 輸出
    W = 112
    print("\n" + "=" * W)
    print(f"{'項目':<26}{'應為值':>12}   報告      簡報      海報      App       判定")
    print("-" * W)
    for key in expected_keys:
        if key in CLAIM_KEY:
            vk = CLAIM_KEY[key]
            nd_, unit_, _eg = CLAIM_FMT.get(vk, (1, "%", ""))
            expv = (half_up(V[vk][0], nd_)[0] + unit_.strip()) if vk in V else "—"
            name = NAMES.get(vk, vk) + "(措辭)"
        else:
            expv = half_up(V[key][0], V[key][1], V[key][2])[0]
            name = NAMES.get(key, key)
        cells, bad, bd, tot = [], 0, 0, 0
        for dv in ("報告", "簡報", "海報", "App"):
            c = stat.get(key, {}).get(dv, [0, 0, 0])
            c2 = stat.get(key, {}).get(dv + "2", [0, 0, 0]) if dv == "App" else [0, 0, 0]
            ok_, bad_, bd_ = c[0] + c2[0], c[1] + c2[1], c[2] + c2[2]
            bad += bad_
            bd += bd_
            tot += ok_ + bad_ + bd_
            cells.append("—" if not (ok_ + bad_ + bd_) else (f"{ok_}✓" + (f"{bad_}★" if bad_ else "") + (f"{bd_}▲" if bd_ else "")))
        verdict = "★不符" if bad else (("○未出現(選用樣式)" if key in OPTIONAL_KEYS
                                        else "○未出現(措辭已改寫,不計)" if key in CLAIM_KEY else "★找不到") if tot == 0
                                       else ("▲邊界" if bd else "OK"))
        print(f"{name[:24]:<26}{expv:>12}   " + "".join(f"{c:<10}" for c in cells) + verdict)
    print("-" * W)
    for fn, loc, msg in app_rows:
        print(f"[App 常數區] {fn} {loc}:{msg}")

    if a.show_ok and ok_list:
        print("\n相符之出現處:")
        for u, key, ctxs, got in ok_list:
            print(f"   ✓ [{u.dv}] {u.loc}  {NAMES.get(key, key)}={got}  …{ctxs}…")
    if a.show_uncovered:
        print("\n同值但未被任何樣式涵蓋之出現處(供人工確認是否需補樣式):")
        for key in expected_keys:
            if key not in V or V[key][1] is None:
                continue
            e = half_up(V[key][0], V[key][1], V[key][2])[0]
            if len(re.sub(r"[^\d]", "", e).lstrip("0")) < 3:
                continue
            rx = re.compile(r"(?<![\d.,])" + re.escape(e) + r"(?![\d])")
            for u in U:
                for m in rx.finditer(u.text):
                    if not any(s0 <= m.start() < s1 for s0, s1 in covered.get((u.dv, u.loc), [])):
                        print(f"   ? [{u.dv}] {u.loc}  {NAMES.get(key, key)}={e}  …{snippet(u.raw, *m.span())}…")
    if exempt:
        print(f"\n○ 依 allowlist.csv 豁免之疑似殘留 {len(exempt)} 處(前後文相符):")
        for dv, loc, name, old, why in exempt:
            print(f"   [{dv}] {loc} {name} {old}:{why[:70]}")

    bad_rows = [r for r in rows if not r["判定"].startswith("▲")]
    bd_rows = [r for r in rows if r["判定"].startswith("▲")]
    print("\n" + "=" * W)
    print(f"逐處比對:相符 {n_ok} 處、不符 {len(bad_rows)} 處(含找不到 {len(never)} 項、無法檢查 {len(underivable)} 項、"
          f"疑似殘留 {len(residual)} 處)、邊界 {len(bd_rows)} 處")
    for r in bad_rows:
        if r["判定"] == "無法檢查":                         # 〔稽核修正 B-10〕依所需檔案彙總列印(工作清單仍逐項)
            continue
        print(f"★ [{r['交付物']}] {r['位置']}|{r['項目']}|現值 {r['現值']} → 應為 {r['應為值']}"
              f"|…{r['前後文']}…" + (f"|{r['備註']}" if r['備註'] else ""))
    if underivable:
        byf = {}
        for k in underivable:
            byf.setdefault(need_file(k), []).append(NAMES.get(k, k))
        for fn, names in byf.items():
            print(f"★ 無法檢查 {len(names)} 項:{fn} 不存在或缺欄(未執行 run_all.py --with-sensitivity?)"
                  f"——{'、'.join(names[:6])}" + (" 等" if len(names) > 6 else ""))
    if optional_absent:
        print(f"○ 選用樣式未出現(不視為錯誤):{'、'.join(NAMES.get(k, k) for k in optional_absent)}")
    for r in bd_rows:
        print(f"▲ [{r['交付物']}] {r['位置']}|{r['項目']}|現值 {r['現值']};候選 {r['應為值']}|來源 {r['來源']}")

    out_csv = a.mismatch_csv or os.path.join(HERE, "crosscheck_mismatch.csv")
    try:
        with open(out_csv, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["交付物", "位置", "項目", "前後文", "現值", "應為值", "判定", "來源", "備註"])
            w.writeheader()
            for r in rows:
                w.writerow(r)
        print(f"\n[工作清單] {len(rows)} 列 → {out_csv}(以 Excel 開啟;學生重跑後,據此逐處更新交付物)")
    except OSError as e:
        print(f"\n[工作清單] 無法寫入 {out_csv}:{e}")

    if bad_rows:
        print("\n結論:未通過 —— 有數字與本次輸出不符、找不到出現處或疑似殘留(見上表與工作清單)。")
        sys.exit(1)
    print("\n結論:通過 —— 樣式涵蓋之每一處數字皆與本次輸出相符"
          + (f"(另有 {len(bd_rows)} 處進位邊界,須以全精度值確認)" if bd_rows else "") + "。")
    print("      注意:本層只檢查樣式涵蓋之位置;樣式外之數字由第 3 層(存在性掃描)與第 4 層補充。")
    sys.exit(0)


if __name__ == "__main__":
    main()
