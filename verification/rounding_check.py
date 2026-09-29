# -*- coding: utf-8 -*-
"""
rounding_check.py —— 第 4 層:交付物表格數值之四捨五入核對
================================================================
把報告、簡報之表格格位(以及海報四象限卡片),回頭與輸出 CSV(或原始資料)之全精度值
比對,確認位數與進位方式正確(ROUND_HALF_UP)。

背景:反向掃描曾逼出一個四捨五入錯誤(報告表 4-6-1 之 RMSE 誤寫 20.59,應為 20.58),
成因為擷取數值時只看了 .round(3) 的顯示值,未回頭核原始浮點。

〔稽核修正 B-17〕(1) 涵蓋率之分子分母改為同口徑:分母＝報告與簡報表格中含數值之格位數,
  分子＝本腳本實際核對之表格格位數;海報四象限卡片為文字框(不在分母內),另列。
  (2) 表 4-1-1 補核最小值/中位數/最大值(原版只核 N、平均、SD、偏態、缺失% 5 欄),且全部
  8 欄改由原始摘要檔 Shanghai_T2DM_Summary.xlsx 以 load_data() 同一轉換全精度計算
  (不再用已捨入至 3 位之 describe.csv)。(3) 另補核表 4-4-1、4-7-1～4-7-5、4-7-7、5-4-1、
  5-4-2、5-4-3 與簡報之馬可夫、CGM 指標表;簡報表格改以表頭內容定位(不再寫死第幾張第幾個圖形)。
〔稽核修正 A-19〕雙重進位邊界:當來源 CSV 已捨入、且其值恰落在目標位數之進位邊界
  (例:CSV 0.8875 → 報告 0.888,全精度 0.887474 應為 0.887)時,無法由 CSV 判定對錯,
  一律列為 ▲ 邊界警告並提示以全精度值確認(不計入失敗)。
〔稽核修正 B-16〕找不到表格(題名改動、表格被刪)或表內變數對不上時,一律計為失敗、
  結束碼非 0(原版只印警告、結束碼 0)。
〔稽核修正 DOCS-17〕交付物與輸出目錄改以標記檔自動搜尋,或以 --base/--out 指定。

[修正 30][修正 36](以題名定位報告表格、象限區間用 quadrant_cost_interval.csv)之設計保留。

用法:
    python rounding_check.py
    python rounding_check.py --base <交付物目錄> --out <程式與輸出目錄>
結束碼:0＝已核對之格位全部正確(邊界警告另列);1＝有不符、找不到表格或無法核對。
"""

# ---------------------------------------------------------------------------
# [修正 32] Windows 主控台編碼防護。
#   本專案之輸出含 ✓ ≤ − ä ² ≈ 等字元,不在繁體中文 Windows 之預設編碼 cp950 內。
#   當 stdout 是「主控台」時 Python 走 WriteConsoleW,不受影響;但當 stdout 被
#   導向「管線」(例如 verify_all.py 以 capture_output=True 抓取子程序輸出,或
#   使用者自行 `python x.py > log.txt`)時,Python 改用地區編碼 cp950 編碼,
#   即拋出 UnicodeEncodeError 並中止 —— 程式本身沒錯,卻因為印不出一個勾勾而失敗。
#   此處只改「遇到無法編碼之字元時的行為」(改為以 ? 取代),不動編碼本身,
#   故主控台顯示維持正常。
import sys as _sys
for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(errors="replace")
    except Exception:
        pass
# ---------------------------------------------------------------------------
import argparse
import decimal
import glob
import os
import re
import sys

import numpy as np
import pandas as pd

D = decimal.Decimal
HERE = os.path.dirname(os.path.abspath(__file__))

# ======================================================================
# 〔稽核修正 DOCS-17〕目錄搜尋
# ======================================================================
_SKIP = {"__pycache__", ".git", "Shanghai_T2DM", "figures_v2", "reference", "reference_teacher",
         "eda_figures", "cgm_output", "output", "output_reg", "output_risk", "output_cost_mc"}


def _is_deliv(d):
    return bool(glob.glob(os.path.join(d, "專題書面報告_v*.docx")))


def _is_out(d):
    return os.path.isfile(os.path.join(d, "output_risk", "quadrant_summary.csv"))


def _find(pred, explicit=""):
    if explicit:
        p = os.path.abspath(explicit)
        return p if pred(p) else ""
    d = HERE
    for _ in range(5):
        frontier, budget = [(d, 0)], 4000
        while frontier:
            x, dep = frontier.pop(0)
            budget -= 1
            if budget < 0:
                break
            if pred(x):
                return os.path.abspath(x)
            if dep >= 4:
                continue
            try:
                subs = sorted(e.path for e in os.scandir(x) if e.is_dir() and e.name not in _SKIP
                              and not e.name.startswith("reference_old_"))
            except OSError:
                continue
            frontier.extend((s, dep + 1) for s in subs)
        parent = os.path.dirname(d)
        if parent == d:
            break
        d = parent
    return ""


# [修正 31] 交付物檔名帶版本號,取版本號最大者。
def pick(base, pattern):
    cs = sorted(glob.glob(os.path.join(base, pattern)))
    if not cs:
        return ""

    def ver(p):
        m = re.search(r"_v(\d+)", os.path.basename(p))
        return int(m.group(1)) if m else 0
    return max(cs, key=ver)


def deliverables(base):
    return {"報告": pick(base, "專題書面報告_v*.docx"),
            "簡報": pick(base, "專題簡報_v*.pptx"),
            "海報": pick(base, "專題海報_80x100cm_v*.pptx")}


# ======================================================================
# 數字工具
# ======================================================================
_TRANS = str.maketrans({"−": "-", "－": "-", "％": "%", "．": ".", "＋": "+"})
_NUM = re.compile(r"(?<![\d.])[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?(?![\d])|(?<![\d.,])[-+]?\d+(?:\.\d+)?(?![\d])")


def nums(text):
    return [m.group(0) for m in _NUM.finditer((text or "").translate(_TRANS))]


def ndec(t):
    return len(t.split(".")[1]) if "." in t else 0


def half_up(v, nd):
    """(ROUND_HALF_UP 字串, 邊界候選或 None)。v 恰有 nd+1 位小數且末位為 5 時為邊界。"""
    d = D(repr(float(v)))
    q = D(1).scaleb(-nd)
    s = str(d.quantize(q, rounding=decimal.ROUND_HALF_UP))
    t = repr(float(v))
    if "e" not in t and "." in t:
        dec = t.split(".")[1]
        if len(dec) == nd + 1 and dec.endswith("5"):
            return s, (str(d.quantize(q, rounding=decimal.ROUND_DOWN)), s)
    return s, None


# ======================================================================
# 原始摘要檔 → 表 4-1-1(與 diabetes_deterioration_pipeline.load_data() 之 d_raw 相同轉換)
# ======================================================================
T411_COLS = {
    "Age": "Age (years)", "BMI": "BMI (kg/m2)", "Duration": "Duration of diabetes (years)",
    "FPG": "Fasting Plasma Glucose (mg/dl)", "PPG": "2-hour Postprandial Plasma Glucose (mg/dl)",
    "HbA1c": "HbA1c (mmol/mol)", "GA": "Glycated Albumin (%)",
    "Cpep_f": "Fasting C-peptide (nmol/L)", "Cpep_2h": "2-hour Postprandial C-peptide (nmol/L)",
    "Ins_f": "Fasting Insulin (pmol/L)", "Ins_2h": "2-hour Postprandial insulin (pmol/L)",
    "TC": "Total Cholesterol (mmol/L)", "TG": "Triglyceride (mmol/L)",
    "HDL": "High-Density Lipoprotein Cholesterol (mmol/L)",
    "LDL": "Low-Density Lipoprotein Cholesterol (mmol/L)", "Cr": "Creatinine (umol/L)",
    "eGFR": "Estimated Glomerular Filtration Rate  (ml/min/1.73m2) ",
    "UA": "Uric Acid (mmol/L)", "BUN": "Blood Urea Nitrogen (mmol/L)",
}


def raw_stats(out_dir):
    p = os.path.join(out_dir, "Shanghai_T2DM_Summary.xlsx")
    if not os.path.exists(p):
        return None
    raw = pd.read_excel(p, sheet_name="T2DM")
    out = {}
    for k, col in T411_COLS.items():
        if col not in raw.columns:
            continue
        x = pd.to_numeric(raw[col], errors="coerce")        # 「/」等非數值 → NaN(同 load_data)
        out[k] = {"N": x.count(), "平均數": x.mean(), "標準差": x.std(), "最小值": x.min(),
                  "中位數": x.median(), "最大值": x.max(), "偏態": x.skew(), "缺失%": x.isna().mean() * 100}
    return out


# ======================================================================
def main():
    ap = argparse.ArgumentParser(description="第 4 層:表格數值四捨五入核對")
    ap.add_argument("--base", default="")
    ap.add_argument("--out", default="")
    a = ap.parse_args()
    BASE = _find(_is_deliv, a.base)
    OUT = _find(_is_out, a.out)
    if not BASE or not OUT:
        sys.exit(f"[錯誤] 找不到交付物或輸出目錄(BASE={BASE!r} OUT={OUT!r})。請用 --base/--out 指定。")
    print("【第 4 層】表格數值四捨五入核對(ROUND_HALF_UP;來源為全精度 CSV 或原始資料)")
    print(f"[path] 交付物:{BASE}")
    print(f"[path] 程式輸出:{OUT}")
    DV = deliverables(BASE)
    for k, v in DV.items():
        print(f"[file] {k}: {os.path.basename(v) if v else '(找不到)'}")

    def rd(rel):
        p = os.path.join(OUT, rel)
        return pd.read_csv(p) if os.path.exists(p) else None

    bad, warn, missing = [], [], []
    checked_cells = set()          # (出處, 表, 列, 欄) —— 表格格位
    poster_boxes = set()

    def chk(where, tkey, ri, ci, label, text, expected, nds=None, arrow=None, exact=False):
        """text 中依序出現之數字,逐一與 expected(全精度)比對;nds=None 時用該數字自身之位數。
        exact=True 表示來源為原始資料之精確十進位值(不做邊界判定)。"""
        if tkey is not None:
            checked_cells.add((where, tkey, ri, ci))
        got = nums(text)
        if len(got) < len(expected):
            bad.append((where, label, text, "應含 %d 個數字" % len(expected)))
            return
        for k, v in enumerate(expected):
            if v is None or (isinstance(v, float) and not np.isfinite(v)):
                missing.append((where, label, "來源值缺"))
                continue
            g = got[k].lstrip("+").replace(",", "")
            nd = (nds[k] if isinstance(nds, (list, tuple)) else nds) if nds is not None else ndec(g)
            e, bd = half_up(v, nd)
            if exact:
                bd = None
            if bd:
                if g in bd:
                    warn.append((where, label, g, f"{bd[0]} 或 {bd[1]}", repr(float(v))))
                else:
                    bad.append((where, label, g, f"{bd[0]} 或 {bd[1]}"))
            elif g != e and not (g in ("-0", "0") and e in ("-0", "0")) and float(g) != float(e):
                bad.append((where, label, g, e))
        if arrow is not None:
            want = "↑" if arrow else "↓"
            if want not in text:
                bad.append((where, label + " 箭頭", text, want))

    # ------------------------------------------------------------ 報告表格(以題名定位)
    import docx
    from docx.oxml.ns import qn
    doc = docx.Document(DV["報告"])
    body = list(doc.element.body.iterchildren())
    TBL, ti = {}, 0
    for i, ch in enumerate(body):
        if ch.tag != qn("w:tbl"):
            continue
        cap = None
        for j in range(i - 1, max(i - 4, -1), -1):
            if body[j].tag == qn("w:p"):
                t = "".join(x.text or "" for x in body[j].iter(qn("w:t"))).strip()
                if t.startswith("表 "):
                    cap = t
                    break
                if t:
                    break
        if cap:
            TBL.setdefault(cap.split("　")[0].strip(), doc.tables[ti])
        ti += 1

    def rows_of(t):
        out = []
        for r in t.rows:
            seen, cells = set(), []
            for c in r.cells:
                if id(c._tc) not in seen:
                    seen.add(id(c._tc))
                    cells.append(c.text.strip())
            out.append(cells)
        return out

    def T(key):
        """〔稽核修正 B-16〕找不到即記為失敗並回傳 None。"""
        if key not in TBL:
            missing.append(("報告", key, f"找不到題名為 {key} 之表格(現有:{', '.join(sorted(TBL))[:120]}…)"))
            return None
        return rows_of(TBL[key])

    M = rd("output/model_results.csv")
    G = rd("output_reg/reg_results.csv")
    Q = rd("output_risk/quadrant_summary.csv")
    S = rd("output_cost_mc/cost_sensitivity_class_weight.csv")
    A = rd("cgm_output/lstm_vs_baseline_all.csv")
    P = rd("cgm_output/lstm_pooled_vs_baseline.csv")
    RT = rd("output_risk/patient_risk_table.csv")
    QC = rd("output_cost_mc/quadrant_cost_interval.csv")
    QORD = ["A", "B", "C", "D"]

    def mrow(t, m):
        r = M[(M.Target == t) & M.Model.astype(str).str.startswith(m)]
        return r.iloc[0] if len(r) else None

    # 表 4-3-1
    R = T("表 4-3-1")
    if R is not None and M is not None:
        order = [("Any_Complication", "Logistic"), ("Any_Complication", "XGB"),
                 ("Microvascular", "Logistic"), ("Microvascular", "XGB"),
                 ("Macrovascular", "Logistic"), ("Macrovascular", "XGB")]
        cols = ["AUC_rep_median", "AP_rep_median", "Precision_rep_median", "Recall_rep_median", "F1_rep_median"]
        for ri, (t, m) in enumerate(order, 1):
            r = mrow(t, m)
            for ci, c in enumerate(cols, 2):
                chk("報告 表 4-3-1", "4-3-1", ri, ci, f"r{ri}c{ci} {t}/{m}/{c}", R[ri][ci], [r[c]], 3)
    # 表 4-6-1
    R = T("表 4-6-1")
    if R is not None and G is not None:
        for ri, mod in [(1, "Linear"), (2, "XGBoost")]:
            r = G[G.Model.astype(str).str.startswith(mod)].iloc[0]
            for ci, (col, nd) in zip((1, 2, 3, 4), [("MAE_median", 2), ("RMSE_median", 2), ("MAPE_median", 1), ("R2_median", 3)]):
                chk("報告 表 4-6-1", "4-6-1", ri, ci, f"r{ri}c{ci} {mod}/{col}", R[ri][ci], [r[col]], nd)
    # 表 4-7-6
    R = T("表 4-7-6")
    if R is not None and A is not None and P is not None:
        chk("報告 表 4-7-6", "4-7-6", 1, 2, "基準 RMSE 平均", R[1][2], [P.baseline_rmse.mean()], 2)
        chk("報告 表 4-7-6", "4-7-6", 2, 2, "逐人 RMSE 平均", R[2][2], [A.lstm_rmse.mean()], 2)
        chk("報告 表 4-7-6", "4-7-6", 2, 3, "逐人 改善中位數", R[2][3], [A.skill_pct.median()], 1)
        chk("報告 表 4-7-6", "4-7-6", 2, 4, "逐人 勝率", R[2][4], [100 * (A.lstm_rmse < A.baseline_rmse).mean()], 1)
        if "n_train_windows" in P.columns:
            chk("報告 表 4-7-6", "4-7-6", 3, 1, "合併 訓練視窗數", R[3][1], [float(P.n_train_windows.sum())], 0)
        chk("報告 表 4-7-6", "4-7-6", 3, 2, "合併 RMSE 平均", R[3][2], [P.lstm_rmse.mean()], 2)
        chk("報告 表 4-7-6", "4-7-6", 3, 3, "合併 改善中位數", R[3][3], [P.skill_pct.median()], 1)
        chk("報告 表 4-7-6", "4-7-6", 3, 4, "合併 勝率", R[3][4], [100 * (P.lstm_rmse < P.baseline_rmse).mean()], 1)
    # 表 4-8-1
    R = T("表 4-8-1")
    if R is not None and Q is not None:
        for ri in range(1, 5):
            q = Q.iloc[ri - 1]
            for ci, col, nd in [(1, "人數", 0), (2, "平均長期風險", 1), (3, "平均短期風險", 1), (4, "實際併發症比例", 1)]:
                chk("報告 表 4-8-1", "4-8-1", ri, ci, f"r{ri} {col}", R[ri][ci], [float(q[col])], nd)
    # 表 4-9-3([修正 30d] 象限層級區間)
    R = T("表 4-9-3")
    if R is not None and QC is not None:
        GQ = QC.assign(k=QC.quadrant.astype(str).str[0]).set_index("k")
        for ri, k in enumerate(QORD, 1):
            g = GQ.loc[k]
            chk("報告 表 4-9-3", "4-9-3", ri, 1, f"{k} 人數", R[ri][1], [float(g["n"])], 0)
            chk("報告 表 4-9-3", "4-9-3", ri, 2, f"{k} 平均(萬)", R[ri][2], [g["mean"] / 1e4], 1)
            chk("報告 表 4-9-3", "4-9-3", ri, 3, f"{k} 區間(萬)", R[ri][3], [g["lo"] / 1e4, g["hi"] / 1e4], 1)
    # 表 4-9-4
    R = T("表 4-9-4")
    if R is not None and S is not None:
        b = S[S.class_weight.astype(str) == "balanced"].iloc[0]
        n = S[S.class_weight.isna()].iloc[0]
        chk("報告 表 4-9-4", "4-9-4", 1, 1, "盛行率 大血管", R[1][1], [b.macro_prev], 1)
        chk("報告 表 4-9-4", "4-9-4", 1, 2, "盛行率 小血管", R[1][2], [b.micro_prev], 1)
        for ri, s in [(2, b), (3, n)]:
            chk("報告 表 4-9-4", "4-9-4", ri, 1, f"r{ri} 大血管預測", R[ri][1], [s.macro_pred], 1)
            chk("報告 表 4-9-4", "4-9-4", ri, 2, f"r{ri} 小血管預測", R[ri][2], [s.micro_pred], 1)
            chk("報告 表 4-9-4", "4-9-4", ri, 3, f"r{ri} AUC", R[ri][3], [s.macro_auc, s.micro_auc], 3)
            chk("報告 表 4-9-4", "4-9-4", ri, 4, f"r{ri} 成本", R[ri][4], [s.cost_mean_first_only / 1e4, s.pop_total / 1e4], [1, 0])
    # 表 4-1-1(原始摘要檔全精度)
    R = T("表 4-1-1")
    RS = raw_stats(OUT)
    if R is not None:
        if RS is None:
            missing.append(("報告", "表 4-1-1", "找不到 Shanghai_T2DM_Summary.xlsx,無法由原始資料核對"))
        else:
            hdr = R[0]
            for ri in range(1, len(R)):
                key = re.split(r"[（(]", R[ri][0])[0].strip()
                if key not in RS:
                    missing.append(("報告 表 4-1-1", f"r{ri}", f"變數 {key!r} 不在原始資料對照表中"))
                    continue
                for ci in range(1, len(hdr)):
                    h = hdr[ci]
                    if h in RS[key]:
                        chk("報告 表 4-1-1", "4-1-1", ri, ci, f"{key} {h}", R[ri][ci], [float(RS[key][h])],
                            exact=h in ("最小值", "中位數", "最大值", "N", "缺失%"))
    # 表 4-4-1(勝算比前六大:特徵、OR、箭頭、名次)
    R = T("表 4-4-1")
    if R is not None:
        tg = {}
        hdr = R[0]
        for ci, h in enumerate(hdr):
            for name, t in (("任一", "Any_Complication"), ("小血管", "Microvascular"), ("大血管", "Macrovascular")):
                if h.startswith(name):
                    tg[ci] = t
        for ci, t in tg.items():
            O = rd(f"output/lr_oddsratio_{t}.csv")
            if O is None:
                missing.append(("報告 表 4-4-1", t, "缺 lr_oddsratio CSV"))
                continue
            O = O.assign(a=O["LR_coef(std)"].abs()).sort_values("a", ascending=False, kind="mergesort").reset_index(drop=True)
            for ri in range(1, min(len(R), 7)):
                feat = re.split(r"[（(]", R[ri][ci])[0].strip()
                checked_cells.add(("報告 表 4-4-1", "4-4-1", ri, ci))
                if feat != str(O.Feature.iloc[ri - 1]):
                    bad.append(("報告 表 4-4-1", f"{t} 第 {ri} 名", feat, str(O.Feature.iloc[ri - 1])))
                    continue
                r = O.iloc[ri - 1]
                chk("報告 表 4-4-1", "4-4-1", ri, ci + 1, f"{t} {feat} OR", R[ri][ci + 1], [r.OddsRatio], 2,
                    arrow=r.OddsRatio > 1)
    # 表 4-7-1(CGM 指標敘述統計;cgm_metrics.csv 全精度,SD 為 ddof=1)
    R = T("表 4-7-1")
    CM = rd("cgm_output/cgm_metrics.csv")
    if R is not None and CM is not None:
        rowmap = {"平均血糖": "mean", "變異係數": "CV", "血糖管理指標": "GMI", "目標範圍內時間": "TIR",
                  "低於範圍時間": "TBR", "高於範圍時間": "TAR"}
        statf = {"平均數": lambda x: x.mean(), "標準差": lambda x: x.std(), "最小值": lambda x: x.min(),
                 "第一四分位": lambda x: x.quantile(.25), "中位數": lambda x: x.median(), "最大值": lambda x: x.max()}
        for ri in range(1, len(R)):
            col = next((v for k, v in rowmap.items() if R[ri][0].startswith(k)), None)
            if col is None or col not in CM.columns:
                missing.append(("報告 表 4-7-1", f"r{ri}", f"列 {R[ri][0]!r} 無法對應 cgm_metrics.csv"))
                continue
            for ci in range(1, len(R[0])):
                f = statf.get(R[0][ci])
                if f:
                    chk("報告 表 4-7-1", "4-7-1", ri, ci, f"{col} {R[0][ci]}", R[ri][ci], [f(CM[col])], 2)
    # 表 4-7-2、4-7-5(轉移矩陣)
    for key, rel, n in (("表 4-7-2", "cgm_output/markov_3state.csv", 3), ("表 4-7-5", "cgm_output/markov_5state.csv", 5)):
        R = T(key)
        MX = rd(rel)
        if R is not None and MX is not None:
            vals = MX.iloc[:, 1:n + 1].to_numpy(dtype=float)
            for i in range(n):
                for j in range(n):
                    chk(f"報告 {key}", key.replace("表 ", ""), i + 1, j + 1, f"[{i},{j}]", R[i + 1][j + 1], [vals[i, j]], 3)
    # 表 4-7-3(穩態、|λ₂|、半衰期;實測時間占比)
    R = T("表 4-7-3")
    MS = rd("cgm_output/markov_stationary_3state.csv")
    if R is not None and MS is not None:
        r = MS.iloc[0]
        cols = list(MS.columns)
        for ci in (1, 2, 3):
            chk("報告 表 4-7-3", "4-7-3", 1, ci, f"穩態 {cols[ci - 1]}", R[1][ci], [r[cols[ci - 1]]], 1)
        chk("報告 表 4-7-3", "4-7-3", 1, 4, "|λ₂|", R[1][4], [r.lambda2], 3)
        chk("報告 表 4-7-3", "4-7-3", 1, 5, "半衰期(小時)", R[1][5], [r.half_life_hours], 1)
        if CM is not None and len(R) > 2:
            for ci, c in ((1, "TBR"), (2, "TIR"), (3, "TAR")):
                chk("報告 表 4-7-3", "4-7-3", 2, ci, f"實測時間占比 {c}", R[2][ci], [CM[c].mean()], 1)
    # 表 4-7-4(收斂過程)
    R = T("表 4-7-4")
    MC = rd("cgm_output/markov_convergence_3state.csv")
    if R is not None and MC is not None:
        mk = {str(k).strip(): i for i, k in enumerate(MC.k_steps.astype(str))}
        cols = list(MC.columns)
        for ri in range(1, len(R)):
            k = R[ri][0].strip()
            if k not in mk:
                if k != "∞":
                    missing.append(("報告 表 4-7-4", f"r{ri}", f"k={k!r} 不在 markov_convergence_3state.csv"))
                continue
            r = MC.iloc[mk[k]]
            for ci, c in ((2, cols[2]), (3, cols[3]), (4, cols[4])):
                chk("報告 表 4-7-4", "4-7-4", ri, ci, f"k={k} {c}", R[ri][ci], [r[c]], 1)
            if nums(R[ri][5]):
                chk("報告 表 4-7-4", "4-7-4", ri, 5, f"k={k} 最大差距", R[ri][5], [r[cols[5]]], 2)
    # 表 4-7-7(依長度與併發症分組)
    R = T("表 4-7-7")
    LG = rd("output/lstm_by_length_group.csv")
    if R is not None and LG is not None and A is not None and P is not None and RT is not None:
        for ri in range(1, 5):
            g = LG.iloc[ri - 1]
            mpt = re.search(r"平均\s*([\d,]+)\s*點", R[ri][0])
            chk("報告 表 4-7-7", "4-7-7", ri, 0, f"長度組 {ri} 平均點數", mpt.group(1) if mpt else R[ri][0],
                [g.mean_points], 0)
            for ci, c in ((1, "per_skill_median"), (2, "pooled_skill_median"), (3, "per_beat_pct"), (4, "pooled_beat_pct")):
                chk("報告 表 4-7-7", "4-7-7", ri, ci, f"長度組 {ri} {c}", R[ri][ci], [g[c]], 1)
        X = (A.merge(P[["record", "lstm_rmse", "skill_pct"]].rename(columns={"lstm_rmse": "pr", "skill_pct": "ps"}), on="record")
              .merge(RT[["record", "has_complication"]], on="record"))
        for ri, comp in ((5, 0), (6, 1)):
            g = X[X.has_complication == comp]
            chk("報告 表 4-7-7", "4-7-7", ri, 0, f"併發症={comp} n", R[ri][0], [float(len(g))], 0)
            for ci, v in ((1, g.skill_pct.median()), (2, g.ps.median()),
                          (3, 100 * (g.lstm_rmse < g.baseline_rmse).mean()), (4, 100 * (g.pr < g.baseline_rmse).mean())):
                chk("報告 表 4-7-7", "4-7-7", ri, ci, f"併發症={comp} c{ci}", R[ri][ci], [v], 1)
    # 表 5-4-1(特徵集敏感度:LR 100 次重複〔稽核修正 A-04,原為 40 次〕;新版 CSV 為全精度,舊版 4 位小數時可能出現進位邊界)
    R = T("表 5-4-1")
    SM = rd("output/sensitivity_missing_features.csv")
    if R is not None and SM is not None:
        tcol = {}
        for ci, h in enumerate(R[0]):
            for name, t in (("任一", "Any_Complication"), ("小血管", "Microvascular"), ("大血管", "Macrovascular")):
                if h.startswith(name):
                    tcol[ci] = t
        for ri in range(1, len(R)):
            try:
                nf = int(nums(R[ri][1])[0])
            except Exception:
                missing.append(("報告 表 5-4-1", f"r{ri}", "無法讀取特徵數"))
                continue
            for ci, t in tcol.items():
                r = SM[(SM.n_features == nf) & (SM.target == t) & SM.model.astype(str).str.startswith("LR")]
                if not len(r):
                    missing.append(("報告 表 5-4-1", f"r{ri}c{ci}", f"CSV 無 {nf} 特徵/{t}/LR"))
                    continue
                r = r.iloc[0]
                chk("報告 表 5-4-1", "5-4-1", ri, ci, f"{nf} 特徵 {t}", R[ri][ci], [r.AUC_median, r.AUC_lo95, r.AUC_hi95], 3)
    # 表 5-4-2、5-4-3(LSTM 時界比較:本研究欄)
    H = rd("cgm_output/lstm_pooled_by_horizon.csv")
    PH60 = rd("cgm_output/lstm_pooled_PH60m.csv")
    PH30 = rd("cgm_output/lstm_pooled_PH30m.csv")
    R = T("表 5-4-2")
    if R is not None and P is not None and PH60 is not None:
        chk("報告 表 5-4-2", "5-4-2", 1, 1, "15 分 本研究 RMSE", R[1][1], [P.lstm_rmse.mean()], 2)
        chk("報告 表 5-4-2", "5-4-2", 2, 1, "60 分 本研究 RMSE", R[2][1], [PH60.lstm_rmse.mean()], 2)
    R = T("表 5-4-3")
    if R is not None and H is not None and PH30 is not None and PH60 is not None:
        h30, h60 = H[H.PH_min == 30].iloc[0], H[H.PH_min == 60].iloc[0]
        for ri, row in enumerate(R):
            lab = row[0]
            if lab.startswith("LSTM RMSE（30"):
                chk("報告 表 5-4-3", "5-4-3", ri, 2, "30 分 彙總/逐紀錄 RMSE", row[2], [h30.pooled_rmse, PH30.lstm_rmse.mean()], 2)
            elif lab.startswith("LSTM RMSE（60"):
                chk("報告 表 5-4-3", "5-4-3", ri, 2, "60 分 彙總/逐紀錄 RMSE", row[2], [h60.pooled_rmse, PH60.lstm_rmse.mean()], 2)
            elif lab.startswith("決定係數"):
                chk("報告 表 5-4-3", "5-4-3", ri, 2, "R² 30/60", row[2], [h30.r2, h60.r2], 3)

    # ------------------------------------------------------------ 簡報表格(以表頭內容定位)
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def flat(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from flat(sh.shapes)
            else:
                yield sh

    pres = Presentation(DV["簡報"])
    slide_tables = []
    for si, s in enumerate(pres.slides, 1):
        for sh in flat(s.shapes):
            if getattr(sh, "has_table", False) and sh.has_table:
                slide_tables.append((si, [[c.text.strip() for c in r.cells] for r in sh.table.rows]))

    def ST(sig):
        # 〔2026-09-29 用詞統一〕sig 之每一項可為字串,或「任一皆可」之多個寫法(tuple)
        alts = [(x,) if isinstance(x, str) else tuple(x) for x in sig]
        for si, rows in slide_tables:
            if all(any(a in h for a in xs for h in rows[0]) for xs in alts):
                return si, rows
        missing.append(("簡報", "/".join(xs[0] for xs in alts), "找不到表頭含此內容之表格"))
        return None, None

    si, R = ST([("邏輯斯迴歸 AUC", "Logistic Regression AUC"), "XGBoost AUC"])
    if R is not None and M is not None:
        for ri in range(1, len(R)):
            t = ("Any_Complication" if "任一" in R[ri][0] else "Macrovascular" if "大血管" in R[ri][0]
                 else "Microvascular" if "小血管" in R[ri][0] else None)
            if t is None:
                continue
            chk(f"簡報 第{si}張", "s-auc", ri, 1, f"{t} LR", R[ri][1], [mrow(t, "Logistic").AUC_rep_median], 3)
            chk(f"簡報 第{si}張", "s-auc", ri, 2, f"{t} XGB", R[ri][2], [mrow(t, "XGB").AUC_rep_median], 3)
    si, R = ST(["合併訓練", "逐人專屬"])
    if R is not None and A is not None and P is not None:
        for ri in range(1, len(R)):
            lab = R[ri][0]
            if lab.startswith(("技巧分數中位數", "改善中位數")):   # 〔2026-09-29 用詞統一〕
                chk(f"簡報 第{si}張", "s-lstm", ri, 1, "合併改善", R[ri][1], [P.skill_pct.median()], 1)
                chk(f"簡報 第{si}張", "s-lstm", ri, 2, "逐人改善", R[ri][2], [A.skill_pct.median()], 1)
            elif lab.startswith("勝過"):
                chk(f"簡報 第{si}張", "s-lstm", ri, 1, "合併勝率", R[ri][1], [100 * (P.lstm_rmse < P.baseline_rmse).mean()], 1)
            elif "RMSE" in lab and "中位數" in lab:
                chk(f"簡報 第{si}張", "s-lstm", ri, 1, "合併 RMSE 中位數", R[ri][1], [P.lstm_rmse.median()], 2)
                chk(f"簡報 第{si}張", "s-lstm", ri, 2, "基準 RMSE 中位數", R[ri][2], [P.baseline_rmse.median()], 2)
    si, R = ST(["平均長期風險", "實際併發症比例"])   # 〔稽核修正 A-24〕簡報第 23 張表頭「人數」已改為「筆數」,不再以此欄名辨識
    if R is not None and Q is not None:
        for ri in range(1, min(5, len(R))):
            q = Q.iloc[ri - 1]
            for ci, col, nd in [(1, "人數", 0), (2, "平均長期風險", 1), (3, "平均短期風險", 1), (4, "實際併發症比例", 1)]:
                chk(f"簡報 第{si}張", "s-quad", ri, ci, f"r{ri} {col}", R[ri][ci], [float(q[col])], nd)
    si, R = ST(["指標（族群中位數）", "數值"])
    if R is not None and CM is not None:
        mp = {"TIR": ("TIR", 1), "TAR": ("TAR", 1), "TBR": ("TBR", 1), "CV": ("CV", 1), "變異係數": ("CV", 1), "GMI": ("GMI", 2)}
        for ri in range(1, len(R)):
            key = next((v for k, v in mp.items() if k in R[ri][0]), None)
            if key:
                chk(f"簡報 第{si}張", "s-cgm", ri, 1, f"{key[0]} 中位數", R[ri][1], [CM[key[0]].median()], key[1])
    si, R = ST(["低血糖", "高血糖"])
    M3 = rd("cgm_output/markov_3state.csv")
    if R is not None and M3 is not None:
        vals = M3.iloc[:, 1:4].to_numpy(dtype=float)
        for i in range(3):
            for j in range(3):
                chk(f"簡報 第{si}張", "s-markov", i + 1, j + 1, f"[{i},{j}]", R[i + 1][j + 1], [vals[i, j]], 3)

    # ------------------------------------------------------------ 海報四象限卡片(文字框,另計)
    po = Presentation(DV["海報"])
    shs = [sh for s in po.slides for sh in flat(s.shapes) if sh.has_text_frame]
    card = {"長高＋短高": "A", "長高＋短低": "B", "長低＋短高": "C", "長低＋短低": "D"}
    if Q is not None and QC is not None:
        GQ = QC.assign(k=QC.quadrant.astype(str).str[0]).set_index("k")
        QQ = Q.assign(k=Q.quadrant.astype(str).str[0]).set_index("k")
        found = 0
        for k, sh in enumerate(shs):
            t = sh.text_frame.text.strip()
            m = re.match(r"^(長[高低]＋短[高低])｜(\d+)\s*筆紀錄$", t)
            if not m or m.group(1) not in card:
                continue
            q = card[m.group(1)]
            found += 1
            poster_boxes.add(sh.name)
            chk("海報 象限卡片", None, 0, 0, f"{q} 人數", m.group(2), [float(QQ.loc[q, "人數"])], 0)
            nxt = shs[k + 1].text_frame.text.strip() if k + 1 < len(shs) else ""
            poster_boxes.add(shs[k + 1].name if k + 1 < len(shs) else "")
            chk("海報 象限卡片", None, 0, 0, f"{q} 併發症/成本", nxt,
                [float(QQ.loc[q, "實際併發症比例"]), GQ.loc[q, "mean"] / 1e4], 1)
        if found != 4:
            missing.append(("海報", "四象限卡片", f"只找到 {found} 張(應為 4)"))

    # ------------------------------------------------------------ 涵蓋率(〔稽核修正 B-17〕同口徑)
    pat = re.compile(r"(?<![\d.])-?\d+(?:\.\d+)?(?![\d])")
    n_rep = sum(1 for t in doc.tables for row in rows_of(t) for c in row if pat.search(c.translate(_TRANS)))
    n_sl = sum(1 for _si, rows in slide_tables for row in rows for c in row if pat.search(c.translate(_TRANS)))
    ck_rep = sum(1 for x in checked_cells if x[0].startswith("報告"))
    ck_sl = sum(1 for x in checked_cells if x[0].startswith("簡報"))
    print(f"\n[涵蓋率] 表格格位(同口徑):報告 {ck_rep}/{n_rep}({ck_rep / max(n_rep, 1) * 100:.1f}%)、"
          f"簡報 {ck_sl}/{n_sl}({ck_sl / max(n_sl, 1) * 100:.1f}%)、"
          f"合計 {ck_rep + ck_sl}/{n_rep + n_sl}({(ck_rep + ck_sl) / max(n_rep + n_sl, 1) * 100:.1f}%)")
    print(f"          另核對海報四象限卡片文字框 {len([b for b in poster_boxes if b])} 個(海報無表格,不在上述分母內)。")
    print("          未核對之格位不代表正確,只代表本腳本沒有查(其中多為文字、文獻值、題號或時步標籤)。")
    covered = sorted({x[1] for x in checked_cells})
    print(f"          已核對之表:{', '.join(covered)}")

    if warn:
        print(f"\n▲ {len(warn)} 格位於進位邊界(來源 CSV 已捨入,恰在 x.xx5 上):以全精度值確認後才能判定")
        for w, l, g, cand, v in warn:
            print(f"   {w} / {l}:表={g!r}  候選={cand}  (CSV 值 {v})")
    if missing:
        print(f"\n★ {len(missing)} 項無法核對(找不到表格或來源,依稽核 B-16 計為失敗):")
        for w, l, why in missing:
            print(f"   {w} / {l}:{why}")
    if bad:
        print(f"\n★ {len(bad)} 項不符:")
        for w, l, g, e in bad:
            print(f"   {w} / {l}:表={g!r}  應={e!r}")
    if bad or missing:
        print("\n結論:未通過。")
        sys.exit(1)
    print("\n結論:已核對之格位全部正確(ROUND_HALF_UP)"
          + (f";另有 {len(warn)} 格進位邊界待以全精度確認。" if warn else "。"))
    sys.exit(0)


if __name__ == "__main__":
    main()
